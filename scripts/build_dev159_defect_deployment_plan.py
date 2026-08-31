from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "CRM_TEST_DEV159_CONSOLIDATED_DEFECT_RCA_AND_DEPLOYMENT_PLAN.md"
OUTPUT = ROOT / "outputs" / "NYSA_CORE_dev159_Consolidated_Defect_Log_RCA_Deployment_and_UAT_Plan.docx"

NAVY = "17324D"
GOLD = "A77B2E"
PALE_GOLD = "F4EBD8"
PALE_BLUE = "EAF0F5"
MID_GREY = "5F6B73"
LIGHT_GREY = "F5F5F2"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=75, start=80, bottom=75, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_runs(paragraph, text, size=9, color=None, bold=False):
    # Minimal Markdown emphasis handling.
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part.replace("`", ""))
        run.bold = bold or (i % 2 == 1)
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("NYSA CORE dev.159 · Local CRM Test candidate · Not deployed · 24 August 2026")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)


def style_document(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(0.55)
    section.footer_distance = Cm(0.6)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string("252A2E")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (("Title", 28, NAVY), ("Heading 1", 19, NAVY), ("Heading 2", 14, GOLD), ("Heading 3", 11.5, NAVY)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def add_masthead(doc):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    left, right = table.rows[0].cells
    set_cell_width(left, 7.8)
    set_cell_width(right, 2.5)
    shade(left, NAVY)
    shade(right, GOLD)
    for cell in (left, right):
        set_cell_margins(cell, 130, 150, 130, 150)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = left.paragraphs[0]
    add_runs(p, "NYSA CORE  |  RELEASE ASSURANCE", 10, WHITE, True)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_runs(p, "DEV.159 · CRM TEST", 10, WHITE, True)


def add_status_banner(doc):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    labels = [
        ("22", "tracked observations"),
        ("1,163 / 0", "ordinary pass / fail"),
        ("22 / 22", "protected DB gates"),
    ]
    for idx, (value, caption) in enumerate(labels):
        cell = table.cell(0, idx)
        set_cell_width(cell, 3.45)
        shade(cell, PALE_GOLD if idx == 0 else PALE_BLUE)
        set_cell_margins(cell, 120, 140, 120, 140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, value, 17, NAVY, True)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, caption, 8.5, MID_GREY)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(set(c) <= {"-", ":", " "} for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_defect_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.autofit = False
    widths = [0.5, 1.15, 1.9, 2.5, 2.35, 2.0]
    header = table.rows[0]
    set_repeat_header(header)
    for idx, text in enumerate(rows[0]):
        cell = header.cells[idx]
        set_cell_width(cell, widths[idx])
        shade(cell, NAVY)
        set_cell_margins(cell, 75, 70, 75, 70)
        p = cell.paragraphs[0]
        add_runs(p, text, 7.5, WHITE, True)
    for row_index, data in enumerate(rows[1:]):
        row = table.add_row()
        prevent_split(row)
        for idx, text in enumerate(data):
            cell = row.cells[idx]
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell, 60, 65, 60, 65)
            shade(cell, WHITE if row_index % 2 == 0 else LIGHT_GREY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            add_runs(p, text, 7.15, NAVY if idx == 0 else None, idx == 0)
    table.style = "Table Grid"


def add_generic_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    set_repeat_header(table.rows[0])
    for idx, text in enumerate(rows[0]):
        cell = table.rows[0].cells[idx]
        shade(cell, NAVY)
        set_cell_margins(cell)
        add_runs(cell.paragraphs[0], text, 8, WHITE, True)
    for ri, data in enumerate(rows[1:]):
        row = table.add_row()
        prevent_split(row)
        for idx, text in enumerate(data):
            cell = row.cells[idx]
            shade(cell, WHITE if ri % 2 == 0 else LIGHT_GREY)
            set_cell_margins(cell)
            add_runs(cell.paragraphs[0], text, 8)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)
    add_masthead(doc)

    title = lines[0].removeprefix("# ")
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_runs(p, title, 27, NAVY, True)
    subtitle = doc.add_paragraph()
    add_runs(subtitle, "Formal defect register · confirmed RCA · correction scope · deployment controls · human UAT package", 12, GOLD, True)

    add_status_banner(doc)
    notice = doc.add_table(rows=1, cols=1)
    cell = notice.cell(0, 0)
    shade(cell, "FFF4E5")
    set_cell_margins(cell, 110, 130, 110, 130)
    add_runs(cell.paragraphs[0], "CONTROL STATUS: Local dev.159 candidate only. No deployment has been performed. Local fixes and automated evidence are not UAT passes.", 9.5, "8A4B08", True)

    i = 1
    first_h2 = True
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows and len(rows[0]) == 6 and rows[0][0] == "ID":
                add_defect_table(doc, rows)
            elif rows:
                add_generic_table(doc, rows)
            continue
        if line.startswith("## "):
            heading = line[3:]
            if heading in {"Consolidated defect, RCA and correction log", "Proposed CRM Test deployment plan — approval required before execution", "Human UAT package — one complete correction round"}:
                doc.add_page_break()
            doc.add_heading(heading, level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("# "):
            pass
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:], 9)
        elif len(line) > 3 and line[0].isdigit() and ". " in line[:4]:
            p = doc.add_paragraph()
            add_runs(p, line, 9)
        elif line.startswith("**") and ":**" in line:
            p = doc.add_paragraph()
            add_runs(p, line, 9)
        else:
            p = doc.add_paragraph()
            add_runs(p, line, 9.2)
        i += 1

    doc.core_properties.title = "NYSA CORE dev.159 Consolidated Defect Log, RCA, Deployment and UAT Plan"
    doc.core_properties.subject = "CRM Test local release candidate assurance"
    doc.core_properties.author = "NYSA CORE Release Assurance"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
