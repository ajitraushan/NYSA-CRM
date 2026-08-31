from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / 'outputs' / 'release_2_6_manuals'
OUT.mkdir(parents=True, exist_ok=True)
NAVY='0B2545'; BLUE='2E74B5'; PALE='E8EEF5'; GRAY='F2F4F7'; RED='9B1C1C'

def set_font(run, size=11, color=None, bold=None, italic=None):
    run.font.name='Calibri'; run._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    run.font.size=Pt(size)
    if color: run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def set_repeat(row):
    trPr=row._tr.get_or_add_trPr(); x=OxmlElement('w:tblHeader'); x.set(qn('w:val'),'true'); trPr.append(x)

def fix_table(table, widths):
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.LEFT
    tblPr=table._tbl.tblPr
    tblW=tblPr.first_child_found_in('w:tblW'); tblW.set(qn('w:w'),'9360'); tblW.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); tblPr.append(ind)
    for row in table.rows:
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margin(cell)

def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    widths=widths or [6.5/len(headers)]*len(headers); fix_table(t,widths)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,PALE); p=c.paragraphs[0]; r=p.add_run(h); set_font(r,10, NAVY, True)
    set_repeat(t.rows[0])
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(str(val)); set_font(r,9.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)
    return t

def setup(doc, label):
    sec=doc.sections[0]; sec.top_margin=Inches(0.8); sec.bottom_margin=Inches(0.75); sec.left_margin=Inches(0.8); sec.right_margin=Inches(0.8); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
    st=doc.styles['Normal']; st.font.name='Calibri'; st._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); st.font.size=Pt(10.5); st.paragraph_format.space_after=Pt(5); st.paragraph_format.line_spacing=1.15
    for name,size,color,before,after in [('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',11.5,NAVY,8,4)]:
        s=doc.styles[name]; s.font.name='Calibri'; s._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=True; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=hp.add_run('NYSA CORE | Release 2.6'); set_font(r,8,NAVY,True)
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=fp.add_run(label+' | Controlled internal manual | Release 2.6 dev.58'); set_font(r,8,'666666')

def cover(doc,title,subtitle):
    doc.add_paragraph().paragraph_format.space_after=Pt(70)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('NYSA REALTY'); set_font(r,12,NAVY,True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(12); r=p.add_run(title); set_font(r,28,NAVY,True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(26); r=p.add_run(subtitle); set_font(r,14,'555555')
    add_table(doc,['Document control','Value'],[['System version','2.1.0-dev.58'],['Release authority','Release 2.6 final UAT / migration 055_release26_transaction_inventory_finance.sql'],['Prepared from','Canonical source, migrations, domain rules, test suite and release runbooks'],['Status','Internal operating reference; features deferred from Release 2.6 are called out explicitly']], [1.7,4.8])
    doc.add_page_break()

def h(doc,t,l=1): doc.add_heading(t,level=l)
def p(doc,t,boldlead=None):
    x=doc.add_paragraph();
    if boldlead:
        r=x.add_run(boldlead); set_font(r,10.5,None,True); x.add_run(t)
    else: x.add_run(t)
    return x
def bullets(doc, items):
    for item in items:
        q=doc.add_paragraph(style='List Bullet'); q.paragraph_format.space_after=Pt(3); q.add_run(item)
def steps(doc,items):
    for item in items:
        q=doc.add_paragraph(style='List Number'); q.paragraph_format.space_after=Pt(3); q.add_run(item)
def note(doc, title, text):
    t=doc.add_table(rows=1,cols=1); fix_table(t,[6.5]); c=t.cell(0,0); shade(c,GRAY); r=c.paragraphs[0].add_run(title+' - '); set_font(r,10,NAVY,True); r=c.paragraphs[0].add_run(text); set_font(r,10); doc.add_paragraph().paragraph_format.space_after=Pt(2)

def user_manual():
    d=Document(); setup(d,'End-User Manual'); cover(d,'NYSA CORE End-User Manual','Role-based CRM, Inventory and Transaction Operations')
    h(d,'1. Purpose, boundaries and operating principles')
    p(d,'This manual describes the workflows implemented in the Release 2.6 canonical source. It is not a policy substitute: users must follow NYSA approval, privacy and evidence standards in addition to the system gates described here.')
    note(d,'Scope boundary','Release 2.6 has governed property-share records and a wa.me handoff, but it does not have live WhatsApp Business/Cloud API delivery callbacks, native property cards/carousels, or production number onboarding. Market-intelligence reporting is future scope, not an available function.')
    h(d,'2. Roles and workspace access')
    add_table(d,['Role','Implemented working scope'],[
      ['Administrator','Company-wide administration, users, teams, controlled values, provider mappings, approvals, dashboards and audit-led governance.'],
      ['Admin Assistant','Operational administration for non-privileged users and teams; cannot alter privileged roles or access status.'],
      ['Managing Director / Director','Company-level dashboard and reporting; proposal and deal approval authority; read-oriented CRM posture.'],
      ['Manager','Managed-team dashboard, lead assignment/reassignment, duplicate and KYC review, Inventory/media review, reservations and deal approval within maintained team scope.'],
      ['Sales Agent','Owned CRM cases: Customers, Leads, qualification, Opportunities, activities, viewings, offers and permitted transaction work.'],
      ['Listing Executive','Inventory workspace and submitted Inventory workflow; no Leads, Customers, Opportunities or My Diary tab.'],
      ['Accountant','Read-oriented access; no My Diary; cannot create CRM records.'],
      ['Partner Broker / Viewer','Restricted collaboration or viewing identity; they do not receive internal Customer CRM access.']], [1.75,4.75])
    h(d,'3. Sign-in, navigation and daily work')
    p(d,'Use the invitation/activation or approved account route, then sign in with your approved credentials. Sessions are cookie-based. Log out on shared devices. Navigation is role-scoped: Dashboard, My Diary, Leads, Opportunities, Customers, Inventory, Reports and Administration appear only when applicable.')
    h(d,'Dashboard and My Diary',2)
    bullets(d,['Dashboard is role-scoped. Directors and Administrators see company executive, sales, inventory, operations and risk views; Managers see maintained-team flow; agents see their operating sequence and assigned priority cases.','Use dashboard counts to open the contributing scoped records. Saved views and scoped CSV export are available in the dashboard workspace.','My Diary brings activities, reminders, meetings and viewings into the user’s working agenda. It is not shown to Listing Executives or Accountants.'])
    h(d,'4. Customers, ownership, duplicates and KYC')
    steps(d,['Search first by maintained name, email and international phone number. A Customer must have a valid email or a valid international phone number.','Choose an existing Customer when a match is correct; otherwise create the Customer with type, communication preference, ownership and available verification/KYC details.','Keep ownership accurate. A Customer owner can maintain the record; an Administrator has broader authority. Avoid creating a new Customer merely to represent an inventory owner - Inventory supports independent seller/landlord/lessor parties.','If a potential duplicate is flagged, do not force a merge. A Manager or Administrator reviews the scoped duplicate queue and records the governed resolution.','Managers/Administrators review KYC in the scoped review queue. Record only supported verification information and evidence; the system’s verification status is not a substitute for external identity-provider verification.'])
    h(d,'Multiple Leads for one Customer',2)
    p(d,'A Customer is the person or organization record; Leads represent separate requirements and lifecycle work. Create another Lead for the same Customer when the requirement, business type, timing, intent or commercial path is distinct. Maintain each Lead’s owner, team, stage, requirements, qualification and activities separately; do not overwrite an earlier Lead to reuse it for a new requirement.')
    h(d,'5. Lead intake, routing, qualification and requirements')
    steps(d,['Create or select the Customer, then create the Lead with source, business type (Sale, Rental, Off-plan or Commercial), objective, team/owner and requirement details.','Route team-queue intake deliberately: the destination team is selected before an eligible active Sales Agent. Managers/Directors/Administrators use the assignment queue within their scope.','Claim or assign promptly. Team response SLA controls the assignment due state; due active Leads appear in the reassignment work.','Capture structured requirements: preferred areas, property attributes, budget range, timing and relevant commercial need. Budget inputs accept normal business notation such as 20K, 2M or AED values.','Complete the approved qualification questionnaire. The calculated result drives Hot/Warm/Cold guidance. An authorized override requires a valid different outcome and a reason.','Move only through the supported Lead flow: New > Contacted > Qualified > Viewing or Negotiation > Won/Lost. A Lost Lead requires a loss reason; only Lost may return to New.'])
    h(d,'6. Opportunity creation and the three governed paths')
    p(d,'Create an Opportunity only from the appropriate qualified starting record. Representation, property source, buyer source, authority evidence and (for dual-sided NYSA) disclosure/conflict evidence are captured and audited.')
    add_table(d,['Path','Create it this way','Key system gate'],[
      ['A. NYSA Customer + NYSA Inventory','Select qualified Customer Lead and approved NYSA Inventory; choose dual-sided NYSA.','Both Lead and approved Inventory are required; both NYSA agents and disclosure evidence are required.'],
      ['B. NYSA buyer + external/co-broker property','Start from qualified Customer Lead; choose buyer representation and an approved provisional external/co-broker property.','External property needs identity, address, source and source evidence, then manager/admin verification before opportunity use.'],
      ['C. NYSA Inventory + external buyer/co-broker','Start from approved NYSA Inventory; choose inventory representation and an external buyer/buyer-agent counterparty.','External counterparty name, type, role, source and evidence are required; Inventory-side NYSA agent is required.']], [1.55,2.6,2.35])
    h(d,'Opportunity workspace',2)
    bullets(d,['Use Requirements and Matching to keep the current requirement and candidate property context accurate.','Record property sharing as a governed share: choose up to 10 eligible properties, recipient/type and selection. The system preserves property snapshots, share items and customer/property-level responses; it opens a wa.me handoff rather than a live provider integration.','Use Viewing, Offer and Negotiation workspaces for the connected operational record. Closed Lost remains available from the supported stages with recorded reason/evidence.'])
    h(d,'7. Inventory: internal record, approval and external listing')
    steps(d,['Create and maintain the internal Inventory record with commercial, location, property and availability facts. Listing Executives work in the dedicated Inventory workspace; permitted Admin Assistants, Managers and Administrators can also create.','Record seller, landlord or lessor parties directly on Inventory. Supported party types include person, company, external broker or agency; Customer creation is not forced.','Add mandate/agreement evidence as applicable: listing/leasing mandate, representation, co-broker, commission, authority, marketing and viewing evidence.','Submit Inventory for internal review. Manager/Administrator can approve, request changes, block or restore according to workflow state. Only approved NYSA Inventory can be selected for governed Inventory/dual Opportunities.','Treat external portal Listing publication as a separate optional workflow. Internal Inventory approval does not equal an external Listing, and external-publication transitions are governed separately.'])
    h(d,'8. Viewings, feedback and calendar behavior')
    steps(d,['Create a viewing against the Opportunity/property context with scheduled date and time, attendees and operational details.','Confirm/reschedule/cancel with the appropriate status and reason. Keep the diary/event record aligned; if Google Calendar is connected, a linked event can synchronize or be reconciled.','Record customer feedback and the next action. A scheduled viewing is visible in the Opportunity and My Diary context; changing a viewing does not itself reserve Inventory.'])
    h(d,'9. Offers, negotiation, booking and reservation')
    bullets(d,['Create an offer revision with exact commercial terms and required evidence. Revisions are immutable; acceptance identifies the accepted revision.','An accepted offer may become an explicit reservation when eligible. Create the reservation with exact amount/currency, refundability, start/end dates and receipt/form evidence.','Reservation blocks the linked Inventory in one governed transaction. The maintained Opportunity Manager can release, expire or cancel it with a reason. Legacy Reserved Inventory without a booking must be reconciled explicitly - it is never silently made available.'])
    h(d,'10. Deals, parties, checklist and authoritative closure')
    steps(d,['Create the Deal from the accepted offer with target completion date and, for Commercial, sale/rental mode. The exact accepted offer amount is carried as agreed value.','Verify transaction parties. A party links exactly one Customer/Contact, Company or transaction-only counterparty and must include source evidence. Inventory seller/landlord/lessor is inherited into the Opportunity and seller side where applicable.','Complete or formally waive required completion checklist items. Required party roles vary by deal type: sale/commercial sale buyer + seller; rental/commercial rental tenant + landlord; off-plan buyer + developer.','Obtain Manager/Director closure approval with a decision basis and evidence reference.','Close Won only after authoritative completion: confirmation, final evidence, completion note and non-future actual completion date are required. Close Lost only with controlled reason, explanation, evidence and explicit confirmation; it releases the reservation as part of closure.'])
    h(d,'11. Mortgage, DBR and investment calculations')
    bullets(d,['Mortgage calculator accepts property price, down payment or loan amount, annual rate, term, costs, income and existing monthly debt. It returns principal, monthly payment, total interest/repayment, LTV and upfront cash.','DBR defaults are prudent 47% and regulatory 50%. The result labels prudent, limited buffer, or above regulatory ceiling and shows the monthly existing-debt reduction needed to reach 47%. This is guidance, not lending approval.','Investment calculation returns effective rent, net income, gross yield, net yield and cash-on-cash return from price, rent, costs, vacancy and cash invested.'])
    h(d,'12. Reports, administration and operational hygiene')
    bullets(d,['Use role-scoped reports for pipeline, sources, activity, agents, movement, calls and closure; export only data in your scope.','Administrators maintain users, invitations, teams, reporting lines, roles, controlled values, portal/provider mappings and organization governance. Privileged roles/access changes have restricted authority and audit records.','Use evidence references, reasons and accurate owners throughout. Audit logging records governed actions. Do not share internal Customer data with partner-broker or portal workflows.'])
    return d

def system_manual():
    d=Document(); setup(d,'System Manual'); cover(d,'NYSA CORE System Manual','Release 2.6 Architecture, Governance, Operations and Controls')
    h(d,'1. Release authority and system boundary')
    p(d,'This manual is grounded in the canonical worktree package version 2.1.0-dev.58, the Release 2.6 final UAT runbook (commit 808620f), additive migrations through 055_release26_transaction_inventory_finance.sql, implemented routes/domain rules, and the repository test suite. The final UAT record states 227/227 automated regressions passed.')
    note(d,'Do not overstate scope','The source explicitly defers live WhatsApp Business/Cloud API integration, approved native property cards/carousels, delivery callbacks, production number onboarding, and market-intelligence reporting. The implemented system instead preserves governed property shares and supports a wa.me handoff.')
    h(d,'2. Runtime and deployment model')
    add_table(d,['Component','Implemented design'],[['Runtime','Node.js 22.13+ application, CommonJS wrapper app.cjs launching ESM server entry point.'],['Database','PostgreSQL; numbered SQL migrations are applied before traffic is accepted.'],['Web client','Lightweight browser app under public/, with role-scoped tabs and domain-specific UI modules.'],['Authentication','Invitation/bootstrap account lifecycle and HttpOnly cookie sessions.'],['Files / PDFs','Private evidence/file workflows plus generated proposal, offer and simple PDF support.'],['Calendar','Optional Google Calendar connection, viewing synchronization/reconciliation and downloadable calendar events.']], [1.5,5.0])
    h(d,'3. Logical modules and primary data relationships')
    add_table(d,['Domain','Purpose and critical relationships'],[['Identity & organization','Brokers/users, role assignments, teams, reporting lines, invitations, access status, organization profile and audit events.'],['CRM','Customer/contact/company records, ownership, KYC summaries, Leads, requirements, qualification, activities and assignment queues. One Customer can support multiple distinct Leads.'],['Inventory','Internal property Inventory, media, workflow approval, party/authority/agreement evidence and optional external portal publication.'],['Opportunity','Qualified requirement-to-property work, matching, representation path, external properties/counterparties, property shares, viewings, offers and bookings.'],['Deal','Accepted-offer transaction record, parties, completion checklist, approval and authoritative won/lost closure.'],['Governance','Controlled values, mappings, approval states, audit history, scope enforcement and evidence-bearing transitions.']], [1.45,5.05])
    h(d,'4. Identity, roles, scope and access controls')
    p(d,'System-level roles are admin, internal_broker, partner_broker and viewer. Internal job roles are Administrator, Admin Assistant, Sales Agent, Listing Executive, Manager, Director and Accountant. Role assignments can be team-scoped and have primary/secondary status, lifecycle dates and approval/change reason.')
    bullets(d,['Team is mandatory for Sales Agent, Listing Executive and Manager assignments. A team has a deliberate Manager and lead response SLA.','Admin Assistant access is deliberately narrower: it cannot alter privileged roles or user access status.','Director/Accountant are treated as non-CRM-writer roles in the UI. Listing Executive is isolated to Inventory and does not receive Leads, Customers, Opportunities or My Diary.','Partner brokers remain outside internal Customer CRM access. External broker identities can be represented as transaction parties or interface identities without a granted internal CRM account.','Customer, lead, team, dashboard and approval queries are scope-aware. Server-side route checks enforce permissions; UI visibility is not the authority.'])
    h(d,'5. Data integrity and lifecycle controls')
    add_table(d,['Area','Implemented controls'],[['Customer identity','Valid email or international phone is required. Search/selection avoids duplicate creation; Manager/Admin duplicate review is separately governed.'],['Lead lifecycle','New -> Contacted -> Qualified -> Viewing/Negotiation -> Won/Lost; Lost requires reason; only Lost can return to New.'],['Qualification','Controlled model accepts non-sensitive factors only; calculated Hot/Warm/Cold result can be overridden only by authorized user with reason.'],['Inventory','Internal workflow supports draft/review/approval/correction/block restoration paths; only approved NYSA Inventory is eligible for governed representation.'],['External property','Identity, address, source and evidence required; provisional external property must be approved for Opportunity use.'],['Offer/reservation','Offer revisions are immutable. Reservation requires accepted revision, amount, dates, refundability and evidence; it blocks Inventory atomically.'],['Deal closure','Mandatory parties, accepted terms, reservation, required checklist and Manager/Director approval gate authoritative Closed Won. Closed Lost requires controlled reason, explanation, evidence and confirmation.']], [1.55,4.95])
    h(d,'6. Representation and transaction design')
    p(d,'Release 2.6 formalizes three representations: buyer, inventory and dual. Representation records persist property source, buyer source, authority evidence, disclosure evidence, counterparties, assigned NYSA agents and commercial attribution. The system validates the combination rather than relying on a descriptive label.')
    add_table(d,['Representation','Required relationship rules'],[['Dual / NYSA Customer + NYSA Inventory','Qualified NYSA Lead + approved NYSA Inventory; NYSA buyer and inventory-side agents; authority and disclosure/conflict evidence.'],['Buyer / external property','Qualified NYSA Lead; provisional external/co-broker property; buyer-side NYSA agent and authority evidence.'],['Inventory / external buyer','Approved NYSA Inventory; external buyer or representative counterparty; inventory-side NYSA agent and authority evidence.']], [2.0,4.5])
    h(d,'7. Inventory versus external Listing publication')
    p(d,'Inventory is the internal authoritative maintained property record. Its approval protects Opportunity selection and transaction work. External Listing publication is a separate optional record/workflow, including provider mapping governance and explicit external-publication state transitions. A user must never assume approved Inventory is public on a portal.')
    h(d,'8. Evidence, audit and files')
    bullets(d,['Governed records require evidence references at key points: representation authority, disclosure, external property source, counterparty provenance, mandate/agreement, offer revision, reservation, deal party source, approval, completion and loss.','Audits are written for material lifecycle and governance operations. Status changes which require reasons capture them rather than silently changing state.','Private file workflows support evidence attachments. Generated proposal/offer PDFs are versions of system artifacts; they do not themselves change commercial state without their corresponding governed action.'])
    h(d,'9. Calculations and decision-support controls')
    add_table(d,['Calculator','Method / output boundary'],[['Mortgage & DBR','Amortization formula using price, down payment/loan, annual rate and term; optional costs/income/debt. Defaults: prudent DBR 47%, regulatory DBR 50%; computes DBR band and existing-debt reduction to prudent threshold.'],['ROI','Annual rent less annual costs divided by property price.'],['Investment return','Effective rent after vacancy, net income, gross/net yield and cash-on-cash return.'],['Qualification','Weighted non-sensitive factors normalized to score; Hot/Warm/Cold thresholding and required override reason.']], [1.55,4.95])
    note(d,'Decision-support constraint','These calculations are transparent operational aids. They do not constitute lender credit approval, formal valuation, regulated advice or an external-provider verification.')
    h(d,'10. Operations, monitoring and release checks')
    steps(d,['Confirm runtime health response, one controlled service PID and expected version identity 2.1.0-dev.58 after deployment.','Confirm migration baseline 054_release26_consolidated.sql and target migration 055_release26_transaction_inventory_finance.sql.','Run the dependency-free test suite with npm test in the canonical project runtime; Release 2.6 final UAT records 227/227 passed at its recorded commit.','Execute focused UAT: DBR example; Inventory party/agreement persistence; party inheritance into Opportunity/Deal; all three representation lanes; representation badge/source persistence; version and migration identity.','Do not modify production or the frozen Release 1.1 candidate when applying Release 2.6 test-environment procedures.'])
    h(d,'11. Administration runbook')
    bullets(d,['Bootstrap only an empty database using the protected BOOTSTRAP_KEY process. Thereafter use approved invitations/user maintenance.','Maintain teams before assigning scoped Sales Agents, Listing Executives or Managers. Keep reporting lines current for Manager-to-Director reporting.','Use controlled-value sets and definitions for governed choices. Stable codes are immutable after activation/use; only unused empty draft records can be deleted.','Provider/listing mapping versions are admin-only. Draft versions are edited/tested; activation retires the prior active provider version according to lifecycle rules.','Suspend/revoke with a reason, which invalidates sessions. Do not change one’s own access state through the administrative route.'])
    h(d,'12. Troubleshooting and safe escalation')
    add_table(d,['Symptom','First safe check'],[['Record not visible','Confirm job role/team scope, ownership and selected dashboard/report filters; do not create a duplicate to bypass scope.'],['Cannot advance state','Read the validation message and complete missing evidence/reason/approval prerequisites.'],['Inventory cannot be reserved','Check approval/availability, accepted offer revision, active booking conflict and evidence. Ask the maintained Manager to release/expire/cancel an existing booking.'],['External property unavailable','Confirm its source/evidence are complete and it is approved for the specific Opportunity.'],['Calendar mismatch','Check connection status and viewing status; use reconciliation rather than manually assuming external state.'],['Unexpected version / migration','Stop rollout, compare health/package identity and migration baseline with the Release 2.6 runbook before further change.']], [1.85,4.65])
    h(d,'13. Explicit non-features in this release')
    bullets(d,['No live WhatsApp provider delivery, callback tracking, approved native cards/carousels or production number onboarding.','No autonomous/ungoverned portal listing publication inferred from Inventory approval.','No customer-facing market intelligence report generation or DXBinteract retrieval workflow.','No automated external identity verification provider: the application records governed verification/KYC status and evidence only.'])
    return d

if __name__=='__main__':
    user_manual().save(OUT/'NYSA_CORE_Release_2_6_End_User_Manual.docx')
    system_manual().save(OUT/'NYSA_CORE_Release_2_6_System_Manual.docx')
    print(OUT)
