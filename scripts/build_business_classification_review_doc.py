from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


OUT = Path(r"C:\Users\ajitr\OneDrive\Desktop\nysa-pocket-ledger\canonical-worktree\outputs\NYSA_CORE_Lead_Classification_Business_Rules_Review_v3.docx")

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "202124"
MUTED = "5F6368"
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF4D6"
PALE_RED = "FDECEC"
BORDER = "CBD2D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            cant_split = OxmlElement("w:cantSplit")
            cant_split.set(qn("w:val"), "true")
            tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=11, bold=False, italic=False, color=INK, font="Arial"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_label_value(doc, label, value, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    set_run(p.add_run(label + ": "), bold=True, color=NAVY)
    set_run(p.add_run(value))
    return p


def add_note(doc, title, text, fill=PALE_GOLD):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(title + "\n"), bold=True, color=NAVY)
    set_run(p.add_run(text), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_decision_block(doc, item, new_page=False):
    if new_page:
        doc.add_page_break()
    doc.add_heading(f"{item['id']}. {item['title']}", level=2)
    add_label_value(doc, "Confirmed current implementation", item["current"])
    add_label_value(doc, "Why confirmation is required", item["risk"])
    add_label_value(doc, "Suggestion for consideration", item["suggestion"])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("Owner review options"), bold=True, color=NAVY)
    selected_option = item.get("selected_option")
    for option_index, option in enumerate(item["options"]):
        op = doc.add_paragraph(style="Review Option")
        marker = "[X]" if selected_option == option_index else "[  ]"
        set_run(op.add_run(f"{marker} {option}"), bold=selected_option == option_index, color=INK)
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [2300, 7060])
    rows = [
        ("Decision status", item.get("decision_status", "PENDING - no decision recorded")),
        ("Owner decision", item.get("owner_decision", "")),
        ("Approved values / rule", ""),
        ("Reason and exceptions", ""),
        ("Migration / effective date", ""),
    ]
    for idx, (label, value) in enumerate(rows):
        left, right = table.rows[idx].cells
        set_cell_shading(left, LIGHT)
        set_run(left.paragraphs[0].add_run(label), bold=True, color=NAVY)
        set_run(right.paragraphs[0].add_run(value), italic=(idx > 0 and not value), color=MUTED if idx > 0 and not value else INK)
        left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    st = styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

if "Review Option" not in styles:
    option_style = styles.add_style("Review Option", WD_STYLE_TYPE.PARAGRAPH)
else:
    option_style = styles["Review Option"]
option_style.base_style = normal
option_style.font.name = "Arial"
option_style.font.size = Pt(10.5)
option_style.paragraph_format.left_indent = Inches(0.2)
option_style.paragraph_format.first_line_indent = Inches(-0.05)
option_style.paragraph_format.space_after = Pt(3)
option_style.paragraph_format.line_spacing = 1.1

# Running header and footer: memo_masthead pattern, no border rule.
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run(hp.add_run("NYSA CORE | Business Rule Review"), size=9, bold=True, color=MUTED)
fp = section.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run(fp.add_run("Owner review draft | 21 August 2026"), size=9, color=MUTED)

# First page memo masthead.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(4)
set_run(p.add_run("BUSINESS RULE REVIEW & DECISION PAPER"), size=22, bold=True, color=NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
set_run(p.add_run("Lead Classification, Routing and Downstream Processing"), size=14, color=MUTED)
add_label_value(doc, "System", "NYSA CORE CRM Test")
add_label_value(doc, "Deployment reviewed", "2.1.0-dev.158")
add_label_value(doc, "UAT references", "UAT-062 and UAT-063")
add_label_value(doc, "Document status", "DRAFT FOR OWNER REVIEW - PARTIAL OWNER DECISIONS RECORDED")
add_label_value(doc, "Prepared date", "21 August 2026")

add_note(
    doc,
    "No-assumption control",
    "This paper records confirmed current implementation, owner-confirmed selections and suggestions for unresolved items. It does not infer any decision: only options explicitly confirmed by the owner are marked [X]. Every blank decision field remains unresolved. No application change, data migration or deployment is authorized by this document.",
    PALE_GOLD,
)

doc.add_heading("1. Purpose and scope", level=1)
doc.add_paragraph(
    "The purpose of this paper is to make the implicit rules attached to the Lead 'Business type' field visible for deliberate review. The review covers Lead capture, Customer-role inference, routing, assignment, qualification, Structured Requirements, Opportunity creation, Inventory eligibility, proposal logic, duplicate handling and reporting."
)
doc.add_paragraph(
    "The evidence basis is the exact cumulative dev.158 package and the corresponding local source. No live CRM data, credentials, private owner/contact information or external service was accessed for this review."
)

doc.add_heading("2. Confirmed issue", level=1)
doc.add_paragraph(
    "The current mandatory field offers Sale, Rental, Off-plan and Commercial as mutually exclusive choices. These values are not one coherent dimension: Sale and Rental describe transaction availability, Off-plan describes market/development stage, and Commercial describes property segment."
)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
set_table_geometry(table, [1680, 2400, 2400, 2880])
headers = ["Current value", "Dimension represented", "Example overlap", "Confirmed concern"]
for idx, text in enumerate(headers):
    set_cell_shading(table.rows[0].cells[idx], PALE_BLUE)
    set_run(table.rows[0].cells[idx].paragraphs[0].add_run(text), bold=True, color=NAVY)
set_repeat_table_header(table.rows[0])
for row in (
    ("Sale", "Transaction availability", "Off-plan commercial sale", "Does not identify buyer versus seller"),
    ("Rental", "Transaction availability", "Commercial office rental", "Does not identify tenant versus landlord"),
    ("Off-plan", "Market/development stage", "Residential or commercial sale", "Not an alternative to Sale"),
    ("Commercial", "Property segment", "Commercial sale or rental", "Not an alternative to Sale/Rental"),
):
    cells = table.add_row().cells
    for idx, text in enumerate(row):
        set_run(cells[idx].paragraphs[0].add_run(text), size=9.5)
set_table_geometry(table, [1680, 2400, 2400, 2880])

doc.add_heading("3. Confirmed processing chain", level=1)
for label, text in (
    ("Lead capture", "The selected value is mandatory and stored as the original Lead classification."),
    ("Routing", "For non-self-assigned creation it participates in source/business/area routing-rule selection."),
    ("Qualification", "It selects a business-line-specific active Qualification Version, with an All-business-lines fallback."),
    ("Requirements", "It initially prefills the versioned Structured Requirement business line; later Requirement versions may diverge."),
    ("Opportunity", "The current Requirement can prefill the Opportunity transaction, which remains separately stored."),
    ("Inventory", "The Opportunity/Requirement transaction is compared by exact text with Inventory transaction types; mismatch is a hard exclusion."),
    ("Owner authority", "Exact Rental follows landlord/lessor handling; other current values generally follow seller/developer handling."),
    ("Proposal", "Off-plan changes the indicative purchase timeline, with additional inference from Inventory handover status."),
    ("Website intake", "The field participates in AI routing, Customer-role inference, continuation/duplicate logic and requirement creation."),
    ("Reporting", "The four values are treated as mutually exclusive business lines in hierarchy/reporting."),
):
    add_label_value(doc, label, text)

doc.add_page_break()
doc.add_heading("4. Owner confirmation register", level=1)
doc.add_paragraph(
    "Complete each decision block independently. Selecting or writing an option confirms only that item. Uncompleted items remain pending and must not be inferred from another answer."
)

decisions = [
    {
        "id": "BR-01", "title": "Lead classification structure",
        "current": "Exactly one of Sale, Rental, Off-plan or Commercial is required.",
        "risk": "The values represent different dimensions and cannot express valid combinations.",
        "suggestion": "Replace the single mixed field with Customer objective, market-stage requirement and property-segment requirement. Derive Sale/Rental processing; do not ask for a separate Lead transaction-availability value.",
        "options": ["Retain the existing four-value single field", "Use separate Customer objective, market-stage requirement and property-segment requirement fields; derive the transaction instead of entering it", "Use an owner-defined alternative taxonomy (specify below)"],
        "selected_option": 1,
        "decision_status": "STRUCTURE SELECTED",
        "owner_decision": "Use Customer objective, market-stage requirement and property-segment requirement on the Lead. Do not add a separate transaction-availability Lead field.",
    },
    {
        "id": "BR-02", "title": "Customer objective values",
        "current": "Sale and Rental do not distinguish which side the Customer is on.",
        "risk": "Buyer/seller and tenant/landlord intent remains ambiguous until later processing.",
        "suggestion": "Use explicit customer-facing objectives that identify both direction and transaction side.",
        "options": ["Use: Buy a property; Sell my property; Rent a property; Rent out my property; Not yet confirmed", "Use only buyer-side objectives: Buy a property; Rent a property", "Use an owner-defined objective list (write every permitted value below)"],
        "selected_option": 0,
        "decision_status": "VALUE SET SELECTED",
        "owner_decision": "Buy a property; Sell my property; Rent a property; Rent out my property; Not yet confirmed.",
    },
    {
        "id": "BR-03", "title": "Derived transaction - not a Lead input field",
        "current": "The current Lead Business type mixes transaction, stage and segment, while Inventory separately stores its offering mode.",
        "risk": "A second independently entered Lead transaction field would duplicate Customer objective and could contradict it.",
        "suggestion": "Derive Sale for Buy/Sell objectives and Rental for Rent/Rent-out objectives. Keep For Sale/For Rent/Both only as an Inventory offering attribute.",
        "options": ["Do not create a Lead transaction field; derive Sale/Rental from Customer objective", "Keep a separate manually entered Lead transaction field", "Use an owner-defined derivation rule (specify below)"],
        "selected_option": 0,
        "decision_status": "DERIVATION RULE SELECTED",
        "owner_decision": "Buy a property or Sell my property derives Sale. Rent a property or Rent out my property derives Rental. Inventory offering mode remains an Inventory attribute.",
    },
    {
        "id": "BR-04", "title": "Market-stage values",
        "current": "Off-plan is embedded in the Business type list and may also be inferred from handover data.",
        "risk": "The maintained value and inferred Inventory condition can conflict.",
        "suggestion": "Consider Ready/Secondary, Off-plan, Either and Not yet confirmed. For Buy/Rent this is a Customer requirement; for Sell/Rent out it comes from linked Inventory.",
        "options": ["Approve this value set: Ready/Secondary; Off-plan; Either; Not yet confirmed", "Approve this reduced value set: Ready/Secondary; Off-plan", "Use an owner-defined market-stage list (write every permitted value below)"],
    },
    {
        "id": "BR-05", "title": "Property-segment values",
        "current": "Commercial is embedded in the Business type list; Residential is only implicit.",
        "risk": "Commercial sale/rental and off-plan commercial combinations cannot be represented coherently.",
        "suggestion": "Consider Residential, Commercial, Either and Not yet confirmed. For Buy/Rent this is a Customer requirement; for Sell/Rent out it comes from linked Inventory.",
        "options": ["Approve this value set: Residential; Commercial; Either; Not yet confirmed", "Do not maintain a separate segment value; derive it from Property type", "Use an owner-defined property-segment list (write every permitted value below)"],
    },
    {
        "id": "BR-06", "title": "Customer-role inference from website intake",
        "current": "Rental creates tenant; investment purpose adds buyer and investor; other values generally create buyer.",
        "risk": "Seller and landlord intent is not inferred, and the mixed classification can create the wrong role.",
        "suggestion": "Derive Customer role from explicit Customer objective, with no inference when objective is unconfirmed.",
        "options": ["Approve objective-based role derivation", "Require manual Customer-role confirmation for every website Lead", "Retain current inference (state accepted exceptions below)"],
    },
    {
        "id": "BR-07", "title": "Website AI routing scope",
        "current": "A determined website AI result is accepted only for Sale or Off-plan; Rental or Commercial goes to Manager review.",
        "risk": "The allowed AI outputs encode the mixed taxonomy and treat two valid categories differently.",
        "suggestion": "Require AI to propose the new independent dimensions, retaining human review and no autonomous assignment.",
        "options": ["Approve dimension-based AI suggestions with human review", "Disable AI classification and route from declared values", "Retain Sale/Off-plan-only AI routing"],
    },
    {
        "id": "BR-08", "title": "Default team routing coverage",
        "current": "Default teams/rules exist for Rental, Off-plan and Secondary Sale; Commercial has no default destination.",
        "risk": "Commercial Leads depend on custom maintenance or Company Unassigned fallback.",
        "suggestion": "Define routing from explicit objective, segment, market stage and area combinations after taxonomy approval.",
        "options": ["Approve multidimensional routing", "Route Commercial to a dedicated team regardless of objective", "Use owner-defined routing coverage (specify below)"],
    },
    {
        "id": "BR-09", "title": "Routing precedence",
        "current": "Lowest numeric priority wins; area specificity is only a tie-breaker after priority. Source/business specificity is not generally preferred.",
        "risk": "A generic lower-numbered rule can override a more specific higher-numbered rule.",
        "suggestion": "Define and display an explicit precedence order before numeric priority is applied.",
        "options": ["Specificity first, then numeric priority", "Numeric priority first, retaining current behavior", "Use an owner-defined precedence order (specify below)"],
    },
    {
        "id": "BR-10", "title": "Sales Agent self-assignment",
        "current": "A Sales Agent who manually creates a Lead is automatically assigned that Lead, bypassing the maintained routing destination.",
        "risk": "Team specialization and area/business routing may be bypassed.",
        "suggestion": "Make the bypass an explicit policy choice and show it before Lead creation.",
        "options": ["Allow self-assignment for every Sales Agent-created Lead", "Apply routing to every Lead", "Allow self-assignment only when the agent is eligible for the resulting route"],
    },
    {
        "id": "BR-11", "title": "Team queue versus direct-agent routing",
        "current": "Routing rules select only a team queue; direct agent assignment is prohibited.",
        "risk": "This may or may not match the intended operating model for specialist or named-account Leads.",
        "suggestion": "Retain team-queue governance unless a separately controlled named-agent exception is approved.",
        "options": ["Retain team-queue-only routing", "Permit governed direct-agent rules", "Use an owner-defined exception policy (specify below)"],
    },
    {
        "id": "BR-12", "title": "Primary-area routing fallback",
        "current": "Primary area is optional; when selected its maintained label must appear in Preferred areas. Without it, only All-areas routing applies.",
        "risk": "Multi-area enquiries may be routed generically even when specialist area teams exist.",
        "suggestion": "Apply search-area fields only to Buy/Rent objectives. For Sell/Rent out, use the linked Inventory's maintained property area; do not ask for Preferred areas or Primary routing area.",
        "options": ["Use objective-dependent area handling: search preferences for Buy/Rent; Inventory area for Sell/Rent out", "Keep the same Preferred-area fields for every objective", "Use an owner-defined area and routing policy (specify below)"],
        "selected_option": 0,
        "decision_status": "FIELD-SOURCE RULE SELECTED; ROUTING DESTINATION STILL PENDING",
        "owner_decision": "Buy/Rent may capture preferred and primary search areas. Sell/Rent out must obtain property area from linked Inventory and must not require those buyer-search fields.",
    },
    {
        "id": "BR-13", "title": "Correction and versioning of Lead classification",
        "current": "The original Lead Business type has no ordinary governed edit/version path after capture.",
        "risk": "Later confirmed information can leave Lead classification inconsistent with Requirements and Opportunity.",
        "suggestion": "Add a governed correction event that preserves original value, reason, actor and effective time.",
        "options": ["Approve governed correction/versioning", "Keep original immutable and derive all processing from current Requirement", "Use an owner-defined correction policy (specify below)"],
    },
    {
        "id": "BR-14", "title": "Qualification model authority",
        "current": "Qualification uses the original Lead Business type and blocks when neither an exact nor All-business model is active.",
        "risk": "A corrected Requirement does not change which questionnaire is selected.",
        "suggestion": "Use the current broker-confirmed Requirement dimensions, with an approved general fallback.",
        "options": ["Use current confirmed Requirement", "Continue using original Lead classification", "Use one common qualification model for all Leads"],
    },
    {
        "id": "BR-15", "title": "Requirement and Opportunity divergence",
        "current": "Requirement business line may differ from Lead; Opportunity transaction is separately selectable and can differ again.",
        "risk": "Three stored classifications can disagree without an explicit reconciliation decision.",
        "suggestion": "Define the current confirmed Requirement as the default authority and require reasoned override at Opportunity creation.",
        "options": ["Require Opportunity to match current Requirement", "Allow override with reason and audit", "Allow unrestricted manual selection"],
    },
    {
        "id": "BR-16", "title": "Inventory mismatch treatment",
        "current": "Exact transaction-type mismatch is a hard exclusion; missing transaction type requires clarification.",
        "risk": "The current mixed taxonomy can exclude valid properties for classification rather than operational reasons.",
        "suggestion": "After separating dimensions, decide which are operational hard gates and which affect ranking/explanation only.",
        "options": ["Hard-gate only Sale versus Rental authority; rank stage/segment/preferences", "Treat every dimension as ranking only", "Use an owner-defined hard-gate list (specify below)"],
    },
    {
        "id": "BR-17", "title": "Seller/landlord/developer authority selection",
        "current": "Exact Rental selects landlord/lessor; other allowed values generally take seller/developer handling.",
        "risk": "Commercial rental and off-plan combinations can select the wrong owner/authority path.",
        "suggestion": "Derive counterparty and authority rules from Customer objective plus Inventory offering mode, not segment or market stage.",
        "options": ["Approve objective/transaction-based authority", "Require manual authority-path selection", "Use an owner-defined authority mapping (specify below)"],
    },
    {
        "id": "BR-18", "title": "Duplicate/continuation and reporting treatment",
        "current": "Business type participates in website duplicate identity, and reporting treats all four values as mutually exclusive lines.",
        "risk": "A classification difference may create a separate Lead, while reports compare unlike dimensions.",
        "suggestion": "Confirm Lead-separation rules by Customer objective and report objective, derived transaction, market stage and segment independently.",
        "options": ["Separate active Leads by Customer objective; report all dimensions", "One active Lead per Customer regardless of objective", "Use an owner-defined continuation/reporting policy (specify below)"],
    },
    {
        "id": "BR-19", "title": "Objective-dependent Lead requirement form",
        "current": "The same buyer-oriented budget, Preferred-area, Primary-routing-area and prompted-property controls are shown for every Lead.",
        "risk": "Seller and landlord Leads must duplicate Inventory facts into buyer-search fields, and the prompted-property wording describes the wrong relationship.",
        "suggestion": "For Buy/Rent, capture search requirements. For Sell/Rent out, select the Customer's Inventory and derive property facts; replace the prompted-property label with Property / Inventory being offered.",
        "options": ["Use objective-dependent forms and derive seller/landlord property facts from linked Inventory", "Keep one common form for every objective", "Use an owner-defined conditional form policy (specify below)"],
        "selected_option": 0,
        "decision_status": "FORM-BRANCH RULE SELECTED; MISSING-INVENTORY PATH STILL PENDING",
        "owner_decision": "Buy/Rent asks for preferred areas, property preferences, budget and timeline. Sell/Rent out links Inventory and derives area, segment, property type, market stage and asking-price context. Do not use Property that prompted this enquiry for the seller/landlord path.",
    },
]

for index, decision in enumerate(decisions):
    add_decision_block(doc, decision, new_page=index > 0)

doc.add_page_break()
doc.add_heading("5. Consolidated target model - mixed decision status", level=1)
add_note(doc, "Decision control", "Customer objective and transaction derivation below are owner-selected. Detailed market-stage and property-segment value sets remain pending. No implementation or deployment is authorized by this paper.", PALE_BLUE)

tax = doc.add_table(rows=1, cols=4)
tax.alignment = WD_TABLE_ALIGNMENT.LEFT
set_table_geometry(tax, [1900, 2860, 2000, 2600])
for i, h in enumerate(("Dimension", "Suggested values", "Used for", "Not to be used for")):
    set_cell_shading(tax.rows[0].cells[i], PALE_BLUE)
    set_run(tax.rows[0].cells[i].paragraphs[0].add_run(h), bold=True, color=NAVY)
set_repeat_table_header(tax.rows[0])
for row in (
    ("Lead: Customer objective", "Buy a property; Sell my property; Rent a property; Rent out my property; Not yet confirmed", "Customer intent, form branch and representation preparation", "Property stage or segment"),
    ("System-derived transaction", "Sale from Buy/Sell; Rental from Rent/Rent out", "Downstream transaction path and matching", "Manual Lead input"),
    ("Lead: Market-stage requirement", "Pending: Ready/Secondary; Off-plan; Either; Not yet confirmed", "Buy/Rent preference; Sell/Rent-out value derives from Inventory", "Sale versus Rental authority"),
    ("Lead: Property-segment requirement", "Pending: Residential; Commercial; Either; Not yet confirmed", "Buy/Rent preference; Sell/Rent-out value derives from Inventory", "Market stage"),
    ("Representation", "Buyer/Tenant; Seller/Landlord; Dual", "Opportunity authority and disclosure", "Lead market classification"),
):
    cells = tax.add_row().cells
    for i, text in enumerate(row):
        set_run(cells[i].paragraphs[0].add_run(text), size=9.5)
set_table_geometry(tax, [1900, 2860, 2000, 2600])

doc.add_heading("6. Combination scenarios for owner review", level=1)
doc.add_paragraph("These examples test whether the approved taxonomy can represent common combinations without inference.")
scenarios = doc.add_table(rows=1, cols=5)
scenarios.alignment = WD_TABLE_ALIGNMENT.LEFT
set_table_geometry(scenarios, [1100, 1650, 1700, 1600, 3310])
for i, h in enumerate(("Scenario", "Objective", "Derived transaction", "Stage / segment", "Requirement source")):
    set_cell_shading(scenarios.rows[0].cells[i], LIGHT)
    set_run(scenarios.rows[0].cells[i].paragraphs[0].add_run(h), bold=True, color=NAVY, size=9.5)
set_repeat_table_header(scenarios.rows[0])
for row in (
    ("A", "Buy property", "Sale", "Ready / Residential", "Customer search requirement"),
    ("B", "Rent property", "Rental", "Ready / Residential", "Customer search requirement"),
    ("C", "Sell property", "Sale", "Ready / Commercial", "Linked Inventory"),
    ("D", "Rent out property", "Rental", "Ready / Commercial", "Linked Inventory"),
    ("E", "Buy property", "Sale", "Off-plan / Residential", "Customer search requirement"),
    ("F", "Sell property", "Sale", "Off-plan / Residential", "Linked Inventory"),
    ("G", "Not confirmed", "Not derived", "Not confirmed", "Customer clarification"),
):
    cells = scenarios.add_row().cells
    for i, text in enumerate(row):
        set_run(cells[i].paragraphs[0].add_run(text), size=9.2, italic=(i == 4))
set_table_geometry(scenarios, [1100, 1650, 1700, 1600, 3310])

doc.add_heading("7. Implementation gate after owner decisions", level=1)
for text in (
    "No implementation starts until all required BR decisions are completed or explicitly marked out of scope.",
    "A field-level data dictionary must be approved, including values, nullability, authority source and correction rules.",
    "Routing, qualification, matching, authority, duplicate and reporting rules must be versioned against the approved dimensions.",
    "Historical values must remain visible as audit evidence; ambiguous Off-plan and Commercial records require governed review rather than silent reinterpretation.",
    "Automated tests and one complete human UAT round must cover every approved combination and exception path before release acceptance.",
):
    p = doc.add_paragraph(style="Review Option")
    set_run(p.add_run("[  ] " + text), color=INK)

doc.add_heading("8. Final owner sign-off", level=1)
sign = doc.add_table(rows=6, cols=2)
sign.alignment = WD_TABLE_ALIGNMENT.LEFT
set_table_geometry(sign, [2800, 6560])
for idx, (label, value) in enumerate((
    ("Overall status", "PENDING"),
    ("Approved by", ""),
    ("Role / authority", ""),
    ("Decision date", ""),
    ("Approved exceptions", ""),
    ("Authorized next action", ""),
)):
    set_cell_shading(sign.rows[idx].cells[0], LIGHT)
    set_run(sign.rows[idx].cells[0].paragraphs[0].add_run(label), bold=True, color=NAVY)
    set_run(sign.rows[idx].cells[1].paragraphs[0].add_run(value), color=INK)
set_table_geometry(sign, [2800, 6560])

doc.add_heading("Appendix A - Evidence references", level=1)
doc.add_paragraph("The review was based on these non-private implementation artifacts in the canonical worktree and exact dev.158 package:")
for text in (
    "public/app.js - visible field values and downstream defaults",
    "src/crm-domain.js - maintained Business type enumeration",
    "src/routing-service.js and src/routes/lead-operations.js - routing selection and defaults",
    "src/routes/qualification-finance.js - Qualification Version selection",
    "src/inventory-eligibility-domain.js - Inventory transaction mismatch handling",
    "src/routes/opportunities.js - Opportunity transaction and owner/authority handling",
    "src/website-ai-routing.js and src/routes/website-intake.js - website AI, role and continuation behavior",
    "src/admin-governance.js - proposal timeline handling",
    "release-artifacts/release-3/consolidated/nysa-core-consolidated-crm-test-dev158.zip - exact deployed package; verified SHA-256 5b6a50ec7bbc13fb7583c9247187fd1fdd09c29870525e508fbeaf8bb98bb8ce",
):
    p = doc.add_paragraph(style="Review Option")
    set_run(p.add_run(text), size=9.5)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = "NYSA CORE Lead Classification Business Rules Review"
doc.core_properties.subject = "UAT-062 and UAT-063 owner review and decision paper"
doc.core_properties.author = "NYSA CORE UAT Documentation"
doc.core_properties.keywords = "NYSA CORE, CRM, UAT-062, UAT-063, business rules, lead classification"
doc.save(OUT)
print(str(OUT))
