from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "NYSA_CRM_TEST_DEV208_LINE_BY_LINE_ISSUE_AND_FIX_REVIEW.docx"
NAVY = "14202B"
GOLD = "C59A45"
PALE = "F3F0E8"
CHARCOAL = RGBColor(0x33, 0x33, 0x33)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(cell, top=75, start=90, bottom=75, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")
        tc_mar.append(tag)


issues = [
    ("My Leave returned 500 for Sales Agent", "Retained the governed leave route and corrected the surrounding access/runtime compatibility so the authenticated empty state loads normally.", "Authenticated Sales Agent request: HTTP 200.", "Open My Leave as a Sales Agent and confirm the screen and balance load."),
    ("Listing Executive dashboard showed the wrong workspace and a redundant Open Inventory action", "The role now lands directly in the Inventory workspace; the duplicate launcher is removed.", "UI/source regression passed.", "Sign in as Listing Executive and inspect the first screen."),
    ("Marketing Compliance returned 500", "Removed the obsolete authorization helper and separated Admin configuration access from operational evidence access.", "Authenticated Admin configuration request: HTTP 200.", "Open configuration as Admin and operational records as a permitted business role."),
    ("Market Intelligence wording was unclear", "Reworded the section as an optional DLD sales comparison with plain business labels.", "UI/source regression passed.", "Confirm that the comparison area is understandable to the business user."),
    ("Known pending task did not reliably appear", "My Tasks now has an explicit assignee path independent of unrelated Lead scope, while other records remain scope-controlled.", "Authenticated Manager task request: HTTP 200.", "Confirm the known assigned pending task appears."),
    ("Controlled external evidence wording was over-regulatory", "Replaced the user-facing language with Required transaction documents and plain step wording; the evidence control itself remains unchanged.", "UI/source regression passed.", "Review wording on the Opportunity/Deal screen."),
    ("Deal completion page did not clearly distinguish open and completed work", "Renamed it Closure Steps and clarified returned, approved, Closed Won and Closed Lost states and authority.", "Deal/domain regression passed.", "Complete one synthetic closure workflow."),
    ("Booking text did not clearly connect to the accepted Offer", "The booking amount now states that it comes from the exact accepted Offer revision; the seven-day reservation rule is retained.", "Release 2 regression passed.", "Reserve synthetic Inventory and inspect the wording."),
    ("Inventory reference missing from Proposal/Offer output", "Proposal output now labels the governed Inventory reference; Offer output already retained it.", "PDF/source regression passed.", "Generate both documents and inspect the reference."),
    ("DEF-124: customer-contact action was hidden below an empty history area", "Placed Record customer contact beside the qualification lock and moved conversation history below the existing governed interaction form.", "UI/source regression passed.", "Record one synthetic Customer contact and verify qualification remains governed."),
    ("DEF-125: Admin could not maintain official-document rules", "Separated Admin configuration endpoints from operational evidence endpoints. Admin can maintain definitions/rules but cannot open operational evidence.", "Authenticated Admin: configuration HTTP 200; operational evidence HTTP 403.", "Draft, activate and retire a synthetic document rule."),
    ("Pale/white Opportunity and Deal blocks remained in the dark theme", "Applied the common charcoal/navy surface treatment to panels, fields and state cards without reintroducing teal.", "Theme/source regression passed.", "Inspect representative Opportunity, Deal, Finance and Admin pages."),
    ("RR-015: no governed purchased Customer bulk import", "Added a separate Customer-only workbook, preview and confirmation workflow. It creates or exactly links Customer Master records and creates no Lead.", "Parser, domain and database integration passed.", "Import synthetic email-only, phone-only, exact-match and conflict rows."),
    ("RR-015: no governed purchased Lead bulk import", "Added a separate Lead-only workflow requiring an existing unambiguous Customer. Leads start New, Unassessed, unassigned and queued; optional preferences create a structured requirement version.", "Parser, domain and database integration passed.", "Import one basic and one enriched synthetic Lead."),
    ("RR-015: duplicate, idempotency and privacy controls", "Added stable source-row idempotency, exact/ambiguous matching, private raw evidence, supplier/date/basis provenance, hash-bound confirmation, a batch register and reconciliation CSV.", "Unit, migration and database integration passed.", "Re-submit a committed source row and inspect the reconciliation register."),
    ("RR-015: importer role boundary", "Only Sales Agent and Admin may import. Manager and Managing Director are denied, and the uploader is never auto-assigned a Lead.", "Authenticated Sales Agent/Admin requests return HTTP 200; Manager/Managing Director return HTTP 403.", "Test both allowed roles and both denied leadership roles."),
]

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.55)
section.right_margin = Inches(0.55)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(8.5)
styles["Normal"].font.color.rgb = CHARCOAL
for name in ("Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Aptos Display"
    styles[name].font.color.rgb = RGBColor(0x9C, 0x72, 0x25)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("NYSA REALTY  |  CORE")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(0xC5, 0x9A, 0x45)

title = doc.add_heading("CRM Test DEV208 — Line-by-Line Issue and Fix Review", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("UAT defect log · DEF-124 · DEF-125 · RR-015")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].italic = True

meta = doc.add_table(rows=2, cols=4)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.style = "Table Grid"
metadata = (("Version", "2.1.0-dev.208", "Status", "Implemented and tested; not deployed"),
            ("Target", "CRM Test only", "Production", "Unchanged and not authorized"))
for row, values in zip(meta.rows, metadata):
    for index, value in enumerate(values):
        cell = row.cells[index]
        cell.text = value
        margins(cell)
        if index % 2 == 0:
            shade(cell, NAVY)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

doc.add_paragraph()
table = doc.add_table(rows=1, cols=5)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["No.", "Issue description", "DEV208 fix", "Evidence", "CRM Test UAT check"]
for cell, text in zip(table.rows[0].cells, headers):
    cell.text = text
    shade(cell, NAVY)
    margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

for number, (issue, fix, evidence, human) in enumerate(issues, 1):
    cells = table.add_row().cells
    values = (str(number), issue, fix, evidence, human)
    for cell, value in zip(cells, values):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        margins(cell)
    if number % 2 == 0:
        for cell in cells:
            shade(cell, PALE)
    cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cells[0].paragraphs[0].runs[0].bold = True

widths = [0.35, 2.15, 3.25, 2.15, 2.6]
for row in table.rows:
    for cell, width in zip(row.cells, widths):
        cell.width = Inches(width)

doc.add_heading("Verification and release boundary", level=1)
for text in (
    "Complete repository regression: 1,522 tests; 1,444 passed; 78 intentionally skipped; 0 failed.",
    "Migration 127 and the purchased-data import contract passed against the dedicated synthetic local PostgreSQL fixture.",
    "The DEV208 archive preserves all 126 historical migrations byte-for-byte and adds only migration 127.",
    "No automatic importer authorization or user-role reassignment is included.",
    "CRM Test deployment requires explicit confirmation. Production and the R2 clone are not targeted.",
):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)

doc.add_heading("Required end-to-end human sequence after an approved CRM Test deployment", level=1)
doc.add_paragraph("Document rule → Deal evidence → Closure → Invoice → Receipt → payout calculation → MD batch approval → Accountant batch payment", style="Intense Quote")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("NYSA Realty · DEV208 review copy · 19 September 2026").font.size = Pt(8)

doc.save(OUT)
print(OUT)
