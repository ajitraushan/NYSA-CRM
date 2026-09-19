from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "documents" / "NYSA-Commission-Payout-Policy-MD-Review.docx"
NAVY = "173746"
GOLD = "B58A3B"
PALE = "F4F1EA"
PALE_BLUE = "EEF4F5"
WHITE = "FFFFFF"
BLACK = "000000"
MID = "5F6B70"
LINE = "D9D9D9"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_border(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def keep_table_together(t):
    for row_index, row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = row_index < len(t.rows) - 1

def set_cell_margins(cell, top=90, start=105, bottom=90, end=105):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def font(run, size=9.5, bold=False, color=BLACK):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run

def add_run(p, text, size=9.5, bold=False, color=BLACK):
    return font(p.add_run(text), size, bold, color)

def add_para(doc, text="", style=None, space_after=5, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.keep_with_next = keep
    if text:
        add_run(p, text)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.05
    add_run(p, text)
    return p

def add_number(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    add_run(p, f"{number}.  ", 9.5, False, BLACK)
    add_run(p, text)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(11 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    font(r, 14 if level == 1 else 11, True, BLACK)
    return p

def table(doc, headers, rows, widths=None, font_size=8.3):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    h = t.rows[0]
    set_repeat_table_header(h)
    for i, label in enumerate(headers):
        c = h.cells[i]
        if widths: c.width = Inches(widths[i])
        set_cell_shading(c, NAVY)
        set_cell_border(c)
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_run(p, label, font_size, True, WHITE)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]
            if widths: c.width = Inches(widths[i])
            set_cell_shading(c, WHITE if ridx % 2 == 0 else PALE_BLUE)
            set_cell_border(c)
            set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i > 0 and len(headers) <= 4:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_run(p, str(value), font_size, False, BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.62)
sec.bottom_margin = Inches(0.62)
sec.left_margin = Inches(0.68)
sec.right_margin = Inches(0.68)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(9.5)
styles["Title"].font.name = "Arial"
styles["Title"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Title"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Title"].font.size = Pt(25)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor.from_string(BLACK)
for name in ("Heading 1", "Heading 2"):
    styles[name].font.color.rgb = RGBColor.from_string(BLACK)

header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_run(hp, "NYSA REALTY", 9, True, GOLD)
add_run(hp, "  |  MD REVIEW", 8, True, MID)

footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(fp, "NYSA Realty  |  Commission payout policy  |  Controlled review copy", 7.5, False, MID)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(4)
font(title.add_run("Commission Payout Policy for MD Review"), 25, True, BLACK)
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(14)
add_run(sub, "Proposed business rules for commission eligibility calculation approval and payment", 11, False, MID)

summary = doc.add_table(rows=1, cols=3)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (label, value) in enumerate((("Document status", "For MD review"), ("Policy owner", "Managing Director"), ("System owner", "NYSA CORE Finance"))):
    c = summary.rows[0].cells[i]
    set_cell_shading(c, PALE)
    set_cell_border(c)
    set_cell_margins(c, 110, 130, 110, 130)
    p = c.paragraphs[0]
    add_run(p, label.upper() + "\n", 7.5, True, GOLD)
    add_run(p, value, 9.5, True, BLACK)

heading(doc, "Decision requested", 1)
add_para(doc, "Approve the calculation rules in this document as the single business specification for Agent commission payouts. The governing principle is that NYSA pays commission only after the related gross commission has been credited to the company account. The achieved tier belongs to the Deal executing Agent, and the resulting Agent pool is then divided using the frozen Deal split.")

heading(doc, "Core policy", 1)
add_bullet(doc, "The property sale price is commercial context and is never the payout basis.")
add_bullet(doc, "The full Deal gross commission equals Buyer commission plus Seller commission recorded for the Deal, excluding VAT.")
add_bullet(doc, "Only gross commission actually credited to NYSA's company account, excluding VAT, enters the payout calculation.")
add_bullet(doc, "The achieved tier rate belongs to the Deal executing or servicing Agent.")
add_bullet(doc, "The achieved tier rate creates one eligible Agent pool for the Deal receipt. The pool is then divided between the servicing and originating Agents using the frozen Deal split.")
add_bullet(doc, "The originating Agent's own tier is not applied again to the same Deal receipt.")
add_bullet(doc, "The social media uplift applies only when the executing Agent has a 100 percent Deal share and is recorded as Social Media Active on the company receipt date.")

heading(doc, "Authoritative data and timing", 1)
table(doc, ["Policy element", "Authoritative record", "Rule"], [
    ("Sale price", "Closed Won Deal", "Displayed for context only"),
    ("Buyer commission", "Frozen Deal commission record", "Excludes VAT"),
    ("Seller commission", "Frozen Deal commission record", "Excludes VAT"),
    ("Full Deal gross commission", "Buyer commission plus Seller commission", "Disclosure and reconciliation value"),
    ("Gross commission received", "Confirmed credit to NYSA company account", "Payout and tier basis"),
    ("Receipt date", "Company account realization date", "Determines the settlement quarter"),
    ("Executing Agent", "Frozen Deal attribution", "Owns the achieved tier"),
    ("Agent split", "Frozen Deal split", "Applied after the tier rate"),
    ("Social media status", "Agent status effective on receipt date", "Additional eligibility test only")
], [1.45, 2.25, 3.25])

doc.add_page_break()
heading(doc, "Calculation sequence", 1)
for number, step in enumerate((
    "Identify the executing Agent from the frozen Deal attribution.",
    "Read the gross commission credited to NYSA for the current receipt, excluding VAT.",
    "Add that receipt to the executing Agent's gross commission received during the same calendar quarter.",
    "Resolve the achieved tier using the inclusive upper-boundary rules.",
    "Apply the social media uplift only when the executing Agent has a 100 percent Deal share and all other uplift conditions are met.",
    "Multiply the eligible gross receipt by the achieved tier rate to create the total eligible Agent pool.",
    "Distribute the pool using the frozen originating and servicing Agent split.",
    "Apply any approved tier increase adjustment required by the active calculation method.",
    "Submit one or more calculated rows to the MD as a consolidated approval batch."
), 1): add_number(doc, number, step)

heading(doc, "Payout formula", 2)
p = add_para(doc)
add_run(p, "Eligible Agent pool = Gross commission received by NYSA ex VAT x Achieved executing Agent tier rate", 10, True, NAVY)
p = add_para(doc)
add_run(p, "Individual Agent payment = Eligible Agent pool x Frozen Agent Deal split", 10, True, NAVY)

heading(doc, "Quarterly tier schedule", 1)
table(doc, ["Quarterly gross commission received", "Base payout rate", "Boundary treatment"], [
    ("AED 0 to AED 100,000", "55 percent", "AED 100,000 remains in this tier"),
    ("Above AED 100,000 to AED 200,000", "60 percent", "AED 200,000 remains in this tier"),
    ("Above AED 200,000 to AED 300,000", "65 percent", "AED 300,000 remains in this tier"),
    ("Above AED 300,000 to AED 400,000", "70 percent", "AED 400,000 remains in this tier"),
    ("Above AED 400,000", "75 percent", "No upper limit")
], [2.7, 1.55, 2.7])
add_para(doc, "The quarter resets at the start of each calendar quarter. A receipt belongs to the quarter containing the date on which the money is realized in the NYSA company account. No unused accumulation carries into the next quarter.")

heading(doc, "Calculation method maintained by Admin", 1)
table(doc, ["Method", "Treatment when a higher tier is reached", "Earlier Deal payouts"], [
    ("Option 1 Quarterly achieved rate", "The new rate applies to eligible gross commission received for all applicable Deals settled in that quarter.", "Recalculated through a separate Tier increase adjustment. Each earlier Deal keeps its original Agent split."),
    ("Option 2 Crossing Deal only", "The new rate applies to the Deal receipt that causes the threshold to be crossed and to later eligible receipts.", "Remain unchanged. No retrospective adjustment is made.")
], [1.65, 2.65, 2.65])
add_para(doc, "Admin must activate one method with an effective date. A calculation preserves the policy version, method and effective date used, so later policy changes do not rewrite an approved or paid calculation.")

heading(doc, "Social media uplift", 1)
add_para(doc, "A Social Media Active status adds five percentage points to tiers 1 to 3 only when the executing Agent has a 100 percent Deal share. This condition reflects the purpose of the uplift: rewarding an Agent who sourced the inventory independently and executed the Deal without an internal split.")
table(doc, ["Executing Agent share", "Social media status", "Base tier", "Applied tier"], [
    ("100 percent", "Active", "55 percent", "60 percent"),
    ("100 percent", "Inactive", "55 percent", "55 percent"),
    ("75 percent with 25 percent originator", "Active", "55 percent", "55 percent"),
    ("Any share below 100 percent", "Active or inactive", "Any tier", "No uplift")
], [1.85, 1.55, 1.45, 2.1])
add_para(doc, "The uplift belongs to the executing Agent's tier calculation. It is not applied separately to the originating Agent and is not added after the Agent pool has been split.")

heading(doc, "Worked example partial receipt and split Deal", 1)
add_para(doc, "A property sells for AED 10,000,000. The Deal records Buyer commission of AED 200,000 and Seller commission of AED 0, so the full Deal gross commission is AED 200,000 excluding VAT. NYSA has received only AED 100,000 in its company account.")
table(doc, ["Calculation line", "Amount or rate"], [
    ("Property sale price", "AED 10,000,000"),
    ("Full Deal gross commission", "AED 200,000"),
    ("Gross commission received and eligible now", "AED 100,000"),
    ("Executing Agent achieved tier", "55 percent"),
    ("Total eligible Agent pool", "AED 100,000 x 55 percent = AED 55,000"),
    ("Servicing Agent payment at 75 percent", "AED 55,000 x 75 percent = AED 41,250"),
    ("Originating Agent payment at 25 percent", "AED 55,000 x 25 percent = AED 13,750"),
    ("Total Agent payments", "AED 55,000")
], [3.25, 3.7])
add_para(doc, "The unpaid AED 100,000 balance does not enter the payout calculation. It becomes eligible only after it is credited to NYSA's account. Because the executing Agent's share is 75 percent, the social media uplift does not apply.")

heading(doc, "Worked example full share with social media uplift", 1)
add_para(doc, "NYSA receives AED 80,000 excluding VAT for a Deal executed and originated by the same Agent. The Agent has a 100 percent Deal share and an active social media status on the receipt date. The base first-tier rate of 55 percent increases to 60 percent.")
table(doc, ["Calculation line", "Amount or rate"], [
    ("Gross commission received", "AED 80,000"),
    ("Base tier", "55 percent"),
    ("Social media uplift", "5 percentage points"),
    ("Applied tier", "60 percent"),
    ("Agent payment", "AED 80,000 x 60 percent x 100 percent = AED 48,000")
], [3.25, 3.7])

heading(doc, "Worked example tier boundary", 1)
add_para(doc, "An executing Agent has AED 100,000 of gross commission received in the quarter. The inclusive upper-boundary rule keeps the Agent in the 55 percent tier. The 60 percent tier begins only when the cumulative amount exceeds AED 100,000.")
table(doc, ["Quarter cumulative", "Achieved base tier"], [
    ("AED 100,000.00", "55 percent"),
    ("AED 100,000.01", "60 percent"),
    ("AED 200,000.00", "60 percent"),
    ("AED 200,000.01", "65 percent")
], [3.45, 3.5])

heading(doc, "Approval and payment workflow", 1)
for number, step in enumerate((
    "The Accountant prepares eligible commission-payment rows after the related company receipt has been confirmed.",
    "The system calculates the tier under the executing Agent, creates one Agent pool and applies the frozen Deal split.",
    "The Accountant selects one or more rows and submits them as one consolidated batch.",
    "The MD reviews the batch in a table and may select or deselect individual rows.",
    "The MD approves the selected rows. Deselected rows remain pending and may be included in a later batch.",
    "After approval, the Accountant records one payment confirmation for the approved batch and supplies the bank payment reference and actual payment date."
), 1): add_number(doc, number, step)

heading(doc, "Required calculation sheet disclosure", 1)
add_para(doc, "Each Agent and settlement quarter must have a printable NYSA calculation sheet. The sheet must show enough information to reproduce every payment without consulting source code.")
table(doc, ["Required field", "Purpose"], [
    ("Settlement date and quarter", "Shows timing and reset period"),
    ("Opportunity and Deal reference", "Links the calculation to the transaction"),
    ("Property sold and sale price", "Identifies the underlying property and commercial value"),
    ("Buyer commission and Seller commission", "Explains the full Deal gross commission source"),
    ("Full Deal gross commission ex VAT", "Shows the contractual commission total"),
    ("Gross commission received ex VAT", "Shows the amount currently eligible for payout"),
    ("Quarter cumulative received", "Supports the achieved tier"),
    ("Executing Agent and achieved tier", "Identifies whose performance determined the rate"),
    ("Social media status and uplift decision", "Explains whether the uplift applied"),
    ("Originating and servicing split", "Explains distribution of the Agent pool"),
    ("Tier increase adjustment", "Separately discloses any retrospective amount under Option 1"),
    ("Payment now", "Shows the amount submitted for approval or payment")
], [2.7, 4.25])

doc.add_page_break()
heading(doc, "Controls and audit requirements", 1)
add_bullet(doc, "The Deal commission terms and Agent split must be frozen before payout calculation.")
add_bullet(doc, "A future company receipt date is prohibited.")
add_bullet(doc, "The same bank or finance reference cannot create a duplicate receipt.")
add_bullet(doc, "The same receipt must not be counted twice in the executing Agent's quarterly cumulative amount.")
add_bullet(doc, "The same Deal gross must not be duplicated across the originating and servicing Agents for tier attainment.")
add_bullet(doc, "Each calculation stores the receipt, Deal, executing Agent, tier policy, social media status, split and method versions used.")
add_bullet(doc, "MD approval and Accountant payment confirmation are separate, attributable audit events.")
add_bullet(doc, "Approved or paid historical calculations remain unchanged when a later policy version is activated.")

heading(doc, "MD approval record", 1)
approval_table = table(doc, ["Decision", "Name and date", "Comments"], [
    ("Approved", "", ""),
    ("Approved with amendments", "", ""),
    ("Returned for revision", "", "")
], [1.8, 2.35, 2.8])
keep_table_together(approval_table)
add_para(doc, "Upon approval, this document becomes the central business specification for implementation and UAT of the NYSA commission-payout workflow.", space_after=0)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
