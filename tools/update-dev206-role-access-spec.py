from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "NYSA_CORE_ROLE_ACCESS_CONDITIONS_DEV205_REVIEW.docx"
OUTPUT = ROOT / "docs" / "NYSA_CORE_ROLE_ACCESS_CONDITIONS_AS_BUILT_DEV206.docx"


def set_cell(cell, value):
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


doc = Document(SOURCE)

replacements = {
    "NYSA CORE Role Access Conditions for Review": "NYSA CORE Role Access Conditions As Built",
    "Proposed DEV205 access model  |  Revised 19 September 2026  |  Decision document":
        "DEV205 deployed model and DEV206 correction candidate  |  Revised 19 September 2026  |  As built control document",
    "Change status  No access-control change has been applied or deployed as part of this document. The matrix below is the proposed model for confirmation before implementation.":
        "Change status  DEV205 is deployed to CRM Test. DEV206 is a local correction candidate addressing Administrator configuration references and leave terminology. Production is unchanged. No deployment may be requested until the control-by-control comparison and focused regression evidence are complete.",
    "Admin Assistant\nDelegated administration and leave administration\nSupports approved administrative tasks and administers leave policy, employment configuration, the leave register and leave decisions. Cannot activate other high-risk settings or change Administrator or Director authority.":
        "Admin Assistant\nLeave administration\nMaintains leave policy and employment configuration, opens the leave register and decides leave. Has no general business or system-configuration access and cannot change Administrator or Director authority.",
    "After confirmation, DEV205 should change role checks, navigation, dashboard selection, query scope and API middleware together. It should not change business calculations, ownership rules, payout formulas, invoice logic, data, schema or production configuration. No database migration is expected. CRM Test deployment will still require a dated source package, checksum, rollback package, focused authorization tests and explicit deployment confirmation.":
        "DEV205 changed role checks, navigation, dashboard selection, query scope and API middleware together. DEV206 corrects configuration-reference access and terminology only. It does not change business calculations, ownership rules, payout formulas, invoice logic, data, schema or production configuration. No database migration is required. Any CRM Test deployment requires a dated source package, checksum, rollback package, focused authorization tests, a completed control-by-control comparison and explicit deployment confirmation.",
    "Source basis  CRM Test 2.1.0-dev.204; access logic inspected 17 September 2026; Git commit 1af87ba994599d8de1bab6d37b2005e609d449fe; existing DEV204 working-tree changes preserved.":
        "Source basis  CRM Test currently serves 2.1.0-dev.205; local DEV206 correction candidate based on Git commit 1af87ba994599d8de1bab6d37b2005e609d449fe; existing working-tree changes preserved; production and R2 clone unchanged.",
}

for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text in replacements:
        paragraph.text = replacements[text]

business_assistant = {
    "Executive business dashboard": "No access",
    "Customers and KYC records": "No access",
    "Leads and assignment": "No access",
    "Inventory and listings": "No access",
    "External portal listings": "No access",
    "Opportunities and Deals": "No access",
    "Proposal and Deal approvals": "No access",
    "Marketing compliance operations": "No access",
    "Receivables and invoice records": "No access",
    "Commission-payment preparation": "No access",
    "Leave application and own balance": "Allowed",
    "Leave approval": "Allowed",
}

governance_assistant = {
    "Users, roles, status and invitations": "No access",
    "Teams and reporting structure": "No access",
    "Audit log and security events": "No access",
    "Organization legal, VAT and bank details": "No access",
    "Payout policy, tiers and effective date": "No access",
    "Controlled values and master data": "No access",
    "Provider mappings and integration setup": "No access",
    "Integration health and failure monitoring": "No access",
    "Listing and media approval policy settings": "No access",
    "Leave policy and employment configuration": "Maintain and administer",
    "Deployment and environment status": "No access",
}

for table in doc.tables:
    if not table.rows:
        continue
    header = [cell.text.strip() for cell in table.rows[0].cells]
    if header and header[0] == "Business function":
        for row in table.rows[1:]:
            key = row.cells[0].text.strip()
            if key in business_assistant:
                set_cell(row.cells[2], business_assistant[key])
    if header and header[0] == "Governance function":
        for row in table.rows[1:]:
            key = row.cells[0].text.strip()
            if key in governance_assistant:
                set_cell(row.cells[2], governance_assistant[key])
    if header and header[0] == "Role":
        for row in table.rows[1:]:
            if row.cells[0].text.strip() == "Admin Assistant":
                set_cell(row.cells[1], "Leave administration")
                set_cell(row.cells[2], "Maintains leave policy and employment configuration, opens the leave register and decides leave. No general business or system-configuration access.")
    if header and header[0] == "No":
        for row in table.rows[1:]:
            number = row.cells[0].text.strip()
            if number == "1":
                set_cell(row.cells[2], "Confirmed decision")
                set_cell(row.cells[3], "Implemented and regression tested")
            elif number == "2":
                set_cell(row.cells[2], "Confirmed decision")
                set_cell(row.cells[3], "Implemented; user assignment requires UAT verification")
            elif number == "3":
                set_cell(row.cells[1], "Admin Assistant is a dedicated leave-administration role and has no general system-configuration or business access.")
                set_cell(row.cells[2], "Confirmed least-privilege decision")
                set_cell(row.cells[3], "Implemented in the central role policy")
            elif number in {"5", "6", "8"}:
                set_cell(row.cells[2], "Implemented security condition")
                set_cell(row.cells[3], "Retain and verify in role UAT")

doc.add_page_break()
heading = doc.add_paragraph("DEV206 control comparison and release gate", style="Heading 1")
heading.paragraph_format.keep_with_next = True
intro = doc.add_paragraph(
    "This comparison is required before any deployment request. Each access condition must be classified as implemented, intentionally excluded by the confirmed model, or blocked pending a named decision. A package is not deployment-ready while any control is unclassified or any required test is failing."
)
intro.paragraph_format.space_after = Pt(8)

rows = [
    ("Administrator navigation", "Only Dashboard and Administration", "Implemented", "Server-provided workspace tabs; no business navigation"),
    ("Administrator dashboard", "Governance status only", "Implemented", "Dedicated Administrator dashboard; no MD mapping"),
    ("Administrator business APIs", "Customers, Leads, Inventory, Opportunities, finance and payouts return 403", "Implemented", "Central deny-by-default governed API policy"),
    ("Administrator team configuration", "Teams and staff references remain available without business-record access", "Corrected in DEV206", "Narrow team and staff configuration capabilities"),
    ("Administrator leave access", "No leave register, policy or decision access", "Implemented", "Leave routes require Admin Assistant capabilities"),
    ("Admin Assistant workspace", "My Leave and Leave Administration only", "Implemented", "Central tabs and deny-by-default API policy"),
    ("Admin Assistant leave ownership", "Policy, employment, register and decision", "Implemented", "Capability-based maintenance and decision checks"),
    ("Managing Director", "Company-wide business oversight; no Administration", "Code implemented", "Authenticated CRM Test UAT still required"),
    ("Accountant", "Dashboard, read-only Opportunities, My Leave, Receivables and Commission Payments", "Code implemented", "Authenticated CRM Test UAT still required"),
    ("Manager and Agents", "Documented scoped business access", "No DEV206 code change", "Focused regression required before deployment"),
    ("Auditability", "Consequential configuration and access changes retain actor and reason", "Code retained", "Authenticated mutation evidence required"),
    ("Role assignments", "No automatic reassignment", "Intentional", "Verify MD is Director and leave approver is Admin Assistant in CRM Test"),
    ("Schema and calculations", "No change", "Confirmed", "DEV206 contains no migration or financial-calculation change"),
    ("Production", "No change without exact approval", "Confirmed", "Production and R2 clone remain untouched"),
]

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
headers = ["Control", "Required condition", "Status", "Evidence or remaining action"]
for index, value in enumerate(headers):
    set_cell(table.rows[0].cells[index], value)
    shade(table.rows[0].cells[index], "243447")
    for run in table.rows[0].cells[index].paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
for row_index, values in enumerate(rows, start=1):
    cells = table.add_row().cells
    for index, value in enumerate(values):
        set_cell(cells[index], value)
        if row_index % 2 == 0:
            shade(cells[index], "F3F5F7")

doc.add_paragraph("Mandatory pre deployment sequence", style="Heading 2")
for text in [
    "Freeze and checksum the exact candidate files.",
    "Complete the control-by-control comparison above against the packaged files, not only the working tree.",
    "Run focused authorization, navigation, leave, Administrator configuration and unchanged-role regression tests.",
    "Perform authenticated CRM Test UAT for Administrator, Admin Assistant, Managing Director and Accountant.",
    "Record test evidence, rollback package, migration status, security and privacy impact, and production state.",
    "Only then present the exact version and change description for deployment confirmation.",
]:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)

for section in doc.sections:
    section.top_margin = section.top_margin

doc.core_properties.title = "NYSA CORE Role Access Conditions As Built"
doc.core_properties.subject = "DEV205 deployed role separation and DEV206 correction release gate"
doc.save(OUTPUT)
print(OUTPUT)
