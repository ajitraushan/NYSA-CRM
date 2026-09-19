from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT=Path(__file__).parent/'business-specifications'
OUT.mkdir(exist_ok=True)
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.65);s.bottom_margin=Inches(.65)
s.left_margin=Inches(.7);s.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(11 if name=='Normal' else 24 if name=='Title' else 17 if name=='Heading 1' else 12)
 st.paragraph_format.space_after=Pt(4)
doc.styles['Normal'].paragraph_format.line_spacing=1
doc.core_properties.title='NYSA Performance Management Business Specification'
doc.core_properties.author='NYSA CRM Development'
doc.core_properties.subject='Owner review and approval'
foot=s.footer.paragraphs[0];foot.alignment=2
foot.add_run('NYSA PM BS 001  |  Review version 0.1  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
for r in foot.runs:r.font.size=Pt(9)
def p(t):doc.add_paragraph(t)
def h(t):doc.add_heading(t,2)
def page(t):
 heading=doc.add_heading(t,1)
 heading.paragraph_format.page_break_before=True
 heading.paragraph_format.space_before=Pt(0)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for vals in rows:
  for c,v in zip(t.add_row().cells,vals):c.text=v
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
 t._tbl.tblPr.append(borders)
 for i,row in enumerate(t.rows):
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
  for c,w in zip(row.cells,widths):
   c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'243B4A' if i==0 else 'F1F4F6' if i%2 else 'FFFFFF');pr.append(sh)
   margins=OxmlElement('w:tcMar')
   for e in ['top','left','bottom','right']:
    x=OxmlElement('w:'+e);x.set(qn('w:w'),'90');x.set(qn('w:type'),'dxa');margins.append(x)
   pr.append(margins)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=1
    for r in para.runs:r.font.size=Pt(10);r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
 doc.add_paragraph().paragraph_format.space_after=Pt(0)

doc.add_paragraph('NYSA Performance Management',style='Title')
doc.add_paragraph('Business Specification for Review',style='Subtitle')
p('Document PM BS 001 | Version 0.1 | 4 September 2026\nPrepared by NYSA CRM Development for the business owner')
h('Decision requested')
p('Approve the proposed KPI definitions, recording rules, monthly target authority and attribution model before development. Publication of this document is for review only. It does not authorize implementation, target activation or deployment.')
h('Business purpose')
p('Performance Management will maintain KPI definitions, assign them to teams and people, publish monthly targets and calculate actuals from governed CRM activity. Agents see their own performance, Managers see their maintained teams, and Directors see company performance. Admin and Managers can set targets within their authority.')
p('My Dashboard will show monthly targets and performance. My Tasks will be the common action queue, including the next assigned action on an Opportunity and role-specific reviews. A KPI drill-down may open its contributing records but must not become a second action queue.')
h('Core requirement')
p('Every live KPI must identify its application fields, qualifying event, time basis, owner, exclusions and drill-down records. No user will type an actual KPI total such as calls made or commission collected into Performance Management. Missing capture must be implemented in the operational workflow before the KPI is activated.')
h('What the application supports today')
p('The dev.179 source contains completed customer interaction capture, controlled call outcomes, task records, qualification assessments, Opportunity stages, offers, bookings, closure and Opportunity-linked receipts. It also has an Admin-only dashboard target facility. It does not yet provide the proposed complete KPI master, Manager allocation workflow, event-time attribution or unified Opportunity task queue.')
p('The field mappings in this specification distinguish existing inputs from new calculation or capture requirements. Existing fields alone do not establish that a KPI is production-ready.')
h('Review order')
p('Pages 2 to 3 cover roles and target setup. Pages 4 to 6 define KPI sources. Page 7 specifies attribution and monthly reporting. Pages 8 to 9 define delivery boundaries, acceptance and approval decisions.')

page('Roles and monthly target governance')
table(['Role','Proposed authority'],[
('Admin','Maintain KPI definitions and approved calculation bindings; configure templates; set and publish targets across company, business line, team and agent scopes.'),
('Director','Set company and team objectives; review company performance; authorize reductions to Director-set team targets. No requirement to approve ordinary Agent operational work.'),
('Manager','Set and publish targets for maintained teams and team members. Allocate a Director-set team target without reducing it. Establish team targets where no higher-level objective exists.'),
('Agent','View own published targets and results; record operational evidence through authorized CRM forms; request a correction, but cannot change targets or actuals directly.'),
('Accountant','Record and reconcile finance evidence through existing permissions. No new target-maintenance, sales-operation or payout-approval rights are granted by this module.')],[1.15,5.9])
h('Monthly workflow')
p('1. Admin activates a versioned KPI definition and relevant role or business-line template.\n2. Director, Admin or Manager sets an objective for a calendar month within scope.\n3. Manager allocates additive team targets across eligible agents. The screen shows team total, allocated total and unallocated balance.\n4. The authorized setter publishes the targets. Each assignee receives a visible target statement. Acknowledgement records receipt, not an approval veto.\n5. Actuals refresh from CRM events. Amendments retain prior values and reasons.\n6. At month end, an authorized Manager closes the team review; company review belongs to Director or Admin. Later corrections create a restated version.')
h('Allocation rules')
p('Counts and AED targets can be allocated by amount or percentage, with deterministic rounding. Publishing an over-allocation requires explicit revision of the parent objective; an under-allocation must show the retained team balance and reason. Parent objectives and child targets must not both be added to company totals.')
p('Percentage and time-based KPIs are thresholds, not additive quotas. A team conversion target of 20% is not split into agent percentages. Team results are recomputed from the underlying numerator and denominator.')
h('Publication safeguards')
p('Allow only one active target per stable KPI code, month, scope and business-line combination, pinned to an approved definition version. Reject overlapping duplicate assignments across versions. Published targets are amended by a new version, never overwritten. A Manager cannot reduce an inherited Director objective; the reduction is routed for approval. Agents changing teams must not erase prior targets or achievements.')

page('Master records and screen fields')
table(['Record','Required fields'],[
('KPI definition','Stable code; name; category; description; unit; direction; target or information-only; approved calculation code/version; source binding; qualifying event; time basis; inclusion/exclusion rules; attribution rule; aggregation; supported roles and business lines; effective dates; status.'),
('KPI template','Code; name; role; business line; selected KPI versions; headline display order; effective dates. No default numerical targets presented as industry benchmarks.'),
('Target assignment','Month; KPI version; scope type and maintained scope ID; business line; target or threshold; parent objective; allocation basis; working-day assumption; reason; creator; publisher; published time; revision; status.'),
('Actual result','KPI and target version; period or as-of time; subject scope; numerator; denominator; value; source event IDs; data completeness; calculation version; refresh time; restatement reference.'),
('Monthly review','Period; scope; original and amended targets; result snapshot; evidence links; explanations; reviewed by/at; closure status; subsequent correction history.')],[1.3,5.75])
h('Use existing application masters')
p('Select agents from brokers.id with active employment/role eligibility and teams from teams.id. Use maintained reporting relationships, not typed team names. Reuse classification and business-line codes used by Leads and Opportunities. Months follow Asia/Dubai boundaries; timestamps remain unambiguous in storage. Supported currency is AED initially.')
h('Existing target records')
p('dashboard_targets already contains metric_code, scope_type, scope_id, period_start, period_end, target_value, unit, definition, exception_threshold, threshold_direction, benchmark_source, approval_basis and status. Existing creation and approval routes are Admin-only. These records require reviewed mapping to the new KPI versions and target lifecycle; do not run two competing target systems or silently import conflicting targets.')
h('Screen design')
p('Performance Management contains KPI Master, Templates and Assignments, Monthly Targets and Allocation, and Performance Review. My Dashboard displays the approved headline subset. Target setting and operational evidence capture remain separate screens. A disabled KPI shows its missing dependency; missing data is not shown as zero achievement.')

page('How calls will be recorded and measured')
h('Existing operational capture')
p('The Agent opens the customer-linked Lead interaction form, selects Call, records Inbound or Outbound, selects the controlled outcome, enters subject and discussion details, optionally enters duration in whole minutes, and provides the next action and due date. Save posts to /api/crm/leads/:id/activities. My Diary also provides a Record call outcome entry point. There is no automatic telephone verification established by this workflow.')
table(['Business fact','Existing application field'],[
('One interaction','activities.id; activity_type = Call; lead_id; contact_id'),
('Direction and outcome','direction; contact_outcome_code. Do not use outcome for the contact result: this route stores the next-action label there.'),
('Completion and time','completed_at is set at saving a completed attempt; created_at records entry creation. A distinct actual call occurrence time is not captured by this form.'),
('Duration and evidence','duration_seconds, converted from durationMinutes; subject; details. Duration is optional, so blank must not imply zero-length or failed contact.'),
('Person and context','owner_id; created_by; opportunity_id_snapshot; opportunity_stage_snapshot; lead_stage_snapshot; qualification_snapshot.'),
('Correction and follow-up','voided_at; activity_corrections; next_action_code; next_action_due_at; next_action_owner_id.')],[1.35,5.7])
h('Proposed call KPI definitions')
p('Outbound call attempts: count distinct non-voided, completed Call activities with direction Outbound in the reporting period. This includes no answer, voicemail, invalid contact and unreachable outcomes. A planned call or opened form does not count.')
p('Connected client conversations: count qualifying Call activities whose controlled outcome is requirements_confirmed, substantive_discussion or callback_requested in the current Call form. Inbound and outbound are shown separately. Additional controlled outcomes require an explicit KPI definition version; do not infer success from free text or duration.')
p('Unique clients reached: count distinct contact_id across connected calls within the selected period and scope. Multiple calls to the same client contribute to conversations, but only one unique client. Team unique-client totals are recomputed, not summed from agent totals.')
h('Example and required extensions')
p('An Agent logs three attempts: no answer, substantive discussion, and callback requested. Proposed result: 3 attempts and 2 connected conversations; unique clients depends on their contact IDs. A later voicemail is another attempt, not another conversation.')
p('Before activation, add an actual occurred_at field with entered-at audit, controlled backdating and duplicate/idempotent submission protection. Add explicit performer attribution and team-at-event snapshot. Opportunity and My Tasks entry points must reuse this same capture with explicit Opportunity linkage; do not create a second call register. Manual entries must be labelled agent-recorded. Telephony integration is a separate future scope.')

page('Activity pipeline and conversion catalogue')
p('All definitions below are proposed. Existing inputs are identified precisely; every KPI still needs the new calculation binding, attribution and acceptance tests.')
table(['KPI','Rule and authoritative inputs','Readiness dependency'],[
('PM01 to PM03\nCalls and clients','Attempt, connected conversation and unique client rules on page 4. activities fields provide the evidence.','Extend capture and counting filters before activation.'),
('PM04\nCompleted meetings','Distinct non-voided activities.id where activity_type = Meeting and completed_at exists; use controlled successful outcomes.','Separate occurrence time from logging time; exclude planned meetings.'),
('PM05\nCompleted viewings','Distinct viewings.id with status completed and nonblank feedback; reference opportunity_id, organizer_id, starts_at and ends_at.','Define actual completion timestamp and audit snapshot; scheduled/no-show/cancelled events excluded.'),
('PM06\nFollow-ups on time','Eligible tasks completed by their applicable due_at divided by eligible tasks due in the period. Source tasks.id, assignee_id, status, due_at, completed_at.','Freeze original and approved revised deadlines; link Opportunity actions and prevent generic completion bypass.'),
('PM07\nNewly qualified leads','Distinct Lead first reaching Qualified from lead_stage_history. qualification_assessments.lead_id, assessed_at and final_temperature provide supporting assessment evidence.','A Warm/Hot assessment alone does not imply Qualified. Capture first valid transition and exclusions.'),
('PM08\nNew Opportunities','Distinct opportunities.id created in month; created_at, owner_id, assigned_team_id, transaction_type and lead_id.','Exclude test/void/duplicate cases under a maintained rule; freeze event ownership.'),
('PM09\nActive pipeline','Distinct open opportunities.id grouped by stage, as of now; next_action_code and next_action_due_at identify stale work.','Snapshot historical stages and action ownership. Do not count shortlisted properties as separate Opportunities.'),
('PM10\nPipeline value','One governed commercial amount per Opportunity; offers.current_revision_id -> offer_revisions.amount/currency where meaningful.','Early stages may have no price. Define estimated_value, basis and expected close date; never sum all shortlisted prices.'),
('PM11\nConversion','Lead-to-qualified and qualified-to-Opportunity use a defined entry cohort. Opportunity win rate = Won / (Won + Lost) for the closure period.','Use transition history and distinct IDs; show sample size, open cohort members and N/A for zero denominator.')],[1.12,3.73,2.2])

page('Results collections and service measures')
table(['KPI','Application source and proposed rule'],[
('PM12 Closed transactions','Count distinct opportunities.id reaching Closed Won, supported by deals.status = closed_won and opportunity_id. Report by opportunities.closed_at. Reconcile closure linkage; multiple Deal rows must not multiply one Opportunity.'),
('PM13 Closed transaction value','deals.agreed_value and currency for the authoritative closure linked to each Opportunity. Rental and sales totals remain separate; rental value basis must be labelled. Use closure date, not booking date.'),
('PM14 Closed commission','Use the governed agreed commission basis linked to the closed Opportunity. Existing buyer/seller commission fields and deal_commission_expectation_versions are inputs, not interchangeable actuals. Freeze one agreed company amount at closure. Label this closed agreed commission, not accounting revenue or agent payout.'),
('PM15 Commission collected','opportunity_commission_receipts: finance_opportunity_id, entry_type, amount, currency, received_date; exact reversals offset receipts. Count each original receipt once. Net commission excludes VAT. Show recorded receipts separately from the currently confirmed reconciliation balance.'),
('PM16 Commission overdue','commission_receivable_invoices: due_date, state, commission_cents, vat_cents, total_cents; schedules.opportunity_id; collections and reversals. Unpaid issued invoices past due are overdue. Display net commission and VAT-inclusive customer balance separately.'),
('PM17 Contact timeliness','leads.first_contact_at plus assignment/acceptance timestamps and the applicable SLA policy. Existing first_contact_at does not alone establish the SLA denominator or reset rules. Reassignment requires a preserved SLA cycle.'),
('PM18 Customer service','Customer satisfaction and complaint resolution are future KPIs. Viewing feedback is narrative, not a satisfaction score. A governed survey/complaint record, scale, timestamps and ownership must exist before these KPIs can be enabled.')],[1.55,5.5])
h('Finance counting and ownership')
p('An invoice payment entered in Receivables already creates its linked net Finance Receipt. Do not add both ledger records to collected commission. Proof upload and receipt confirmation are not new cash receipts. Historical unallocated receipts require reviewed allocation before invoice-specific performance is claimed.')
p('Collection attribution belongs to the underlying Opportunity and approved split, not the Accountant who enters the receipt. Team/company totals count company commission once. Agent shares use a frozen originating/servicing attribution version; splits on the Opportunity are percentage inputs, not proof of agent identity or payout entitlement.')
h('Recommended initial headline set')
p('Agents: connected conversations, unique clients reached, completed viewings, new Opportunities, closed transactions and closed agreed commission. Managers add team conversion and overdue actions. Directors focus on company closures, commission, collection against due amounts and pipeline health. Metrics with missing dependencies remain informational or unavailable until verified.')

page('Attribution monthly status and data quality')
h('Attribution rules')
p('Activity KPIs credit the actual performer, not whoever currently owns the Lead. Preserve performer, team and business line at event time. created_by and owner_id must be reconciled where someone records on behalf of another person. Do not assume a current team join can reconstruct historical membership.')
p('For newly created Opportunities, credit the owner/team at creation. For closure count and transaction value, propose one accountable servicing owner/team at closure; display originating participation separately, without duplicating company totals. Monetary agent attribution uses the agreed originating/servicing split frozen at the relevant event. Management approval of these definitions is required.')
h('Monthly calculations')
p('Use calendar months in Asia/Dubai. Event totals such as calls and closures are month-to-date; pipeline is a point-in-time balance. For a positive count or AED target, achievement is actual divided by target, expressed as a percentage. Remaining gap cannot fall below zero. A target of zero or no target yields N/A, not a division error.')
p('Percentage KPIs show actual percentage, target percentage and percentage-point gap. Lower-is-better measures use their own threshold logic. Do not label a short response time as low achievement simply because it is below the target.')
p('Pace uses elapsed approved working days only for additive activity quotas. Do not linearly forecast conversion, pipeline balance or closings by default. A weighted pipeline forecast needs maintained stage probabilities, expected close dates and a clear estimate label; these are future dependencies, not current facts.')
h('Data and review controls')
p('Every result must open the qualifying source records, show exclusions and expose refresh time and calculation version. Missing data, an empty valid population and a genuine zero result are different states. Current month may refresh live; closed-month snapshots must remain reproducible.')
p('Voids, reversals, corrections and late entries recompute the open month. Closed-month changes require an authorized restatement retaining the original published result and reason. The operational record is corrected first; users cannot edit the KPI total itself. Guard against repeated stage transitions, duplicate imports and retried submissions.')
p('Approved leave, new starters and part-month transfers may justify revised activity targets; do not automatically prorate an agreed revenue objective. Show original target, adjustment, effective days and approver. Proposed backdated activity entries require Manager review when the reporting period is closed.')
h('Security and working boundaries')
p('Apply role and team scope on the server and on drill-down/export endpoints. Agents cannot read others through a KPI link; Managers cannot change another team. Do not include phone numbers, private proof contents or salaries in aggregate KPI exports. Performance does not change payout authority, commission receipt status, customer-contact restrictions or transaction closure rules.')

page('Delivery scope and acceptance criteria')
h('Implementation sequence after approval')
p('First establish KPI definitions, source adapters and missing event capture. Then implement role-scoped target assignment and publication, followed by monthly results and drill-downs. Finally separate My Dashboard from My Tasks and reconcile historical target records. Deliver in CRM Test with synthetic fixtures and a recoverable migration plan; production requires separate explicit approval.')
p('My Tasks must collect each assigned operational action once, show its source, due date and exact action link, and remove or advance it only when the governed workflow changes. Manager and Director approval queues may retain dedicated review forms, but their pending assignments must appear in the unified queue. KPI definitions and targets must not themselves generate duplicate operational tasks.')
table(['Test','Required result'],[
('AT01 Call evidence','An outbound no-answer Call adds one attempt and no connected conversation. A connected Call adds one to each. Opening a form or scheduling a call adds neither.'),
('AT02 Corrections','Duplicate retry creates one activity. Void excludes it. A correction/reversal updates current results once and retains source audit evidence.'),
('AT03 Source reconciliation','Every displayed count/amount reconciles to the drill-down IDs. Missing occurrence time or identity is visible, not silently attributed.'),
('AT04 Targets and roles','Manager can publish for own team; Agent cannot edit; cross-team access is rejected. Director-set objective cannot be reduced without authorization.'),
('AT05 Aggregation','Team rates use summed numerators/denominators; unique clients and Opportunities are deduplicated. Parent and child targets are not double-counted.'),
('AT06 Finance','An AED1,050 invoice payment at 5% VAT produces AED1,000 net collected commission once. Proof upload and confirmation add no extra cash.'),
('AT07 Time and transfer','Dubai month boundary, midmonth transfer, leave adjustment and late correction preserve original target and historical attribution.'),
('AT08 Task and dashboard','Current Opportunity next action appears once for the responsible role. Dashboard shows performance, and only the correct governed form completes the work.'),
('AT09 Closure independence','A Manager can close an eligible transaction with unpaid commission. Collection stays outstanding; no automatic payout approval occurs.')],[1.2,5.85])

page('Approval decisions and application traceability')
h('Owner decisions requested')
p('Confirm or amend: (1) the initial headline KPI set; (2) connected-call outcome definitions and manual recording disclosure; (3) performer/team-at-event attribution; (4) Manager publication rights and inherited target reduction approval; (5) servicing-owner attribution for transaction counts and split attribution for commission; (6) closed-month restatement and backdating rules; (7) which capture dependencies must be included in the first release. No numeric industry quotas are assumed.')
p('Review outcome: Approve / Approve with changes / Revise.\nReviewer: ____________________  Date: ____________________\nComments and conditions: __________________________________________')
h('Inspected application baseline')
p('Field mappings were checked against the exact isolated CRM Test dev.179 source package, SHA256 e6b4d2d2dc734b03873aecb6c4e4a5d9d902c739ac70f0d9481d86c347306434. This is a source-based specification, not a fresh live data audit. Existing deployment evidence records migrations through116 and a paired Test rollback at /home/nysareal/crm-backups/consolidated-crm-test-dev179-20260903T191641Z.')
table(['Source file in dev179','Verified responsibility'],[
('public/app.js','Call/Meeting interaction form, controlled outcomes, duration conversion, My Diary entry and current My Tasks renderer.'),
('src/routes/crm.js','Activity creation, completion time, Opportunity snapshot, controlled outcome validation, corrections and voids.'),
('src/routes/dashboards.js','Current broad Call count/report and Admin-only dashboard_targets create/approve routes.'),
('src/routes/lead-operations.js','Current assigned task query; no automatic projection of ongoing Opportunity next actions.'),
('src/routes/qualification-finance.js','qualification_assessments with assessed_at, final_temperature and assessor.'),
('src/routes/opportunities.js','Viewing outcome recording, immutable offer revisions, booking, closure and linked Opportunity state.'),
('src/opportunity-finance.js and src/commission-receivables.js','Opportunity-based receipt and reconciliation queries; invoice schedules and collection operations.'),
('src/migrations/113, 115 and 116 SQL files','Receivables schema, linked collection/receipt posting and effective Opportunity finance views.')],[2.6,4.45])
h('Publication and exclusions')
p('This review version publishes the business proposal only. No CRM source code, live target, database schema, access role or deployment has been changed. Production, R2 clone and Property Finder are excluded. Existing human UAT observations and automated results remain unchanged. Automated telephony, call audio, customer surveys, complaint management, payout redesign and statutory invoice generation are outside this specification unless separately approved.')
for root in [doc.styles.element, doc.element]:
 for border in list(root.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
doc.save(OUT/'NYSA_Performance_Management_Business_Specification_v0.1.docx')
print(OUT/'NYSA_Performance_Management_Business_Specification_v0.1.docx')
